#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import os, zipfile
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_BREAK

BASE = "/home/user/6-main/01-6个主号内容创作/引流小号_发布文案"
ZIPNAME = "编号20260712-01_引流账号文案_芳芳财会.zip"
DOCNAME = "引流小号_第232-240篇.docx"
P_BASE = "#会计 #中级会计 #中级会计备考 #注会cpa #财会"
P_TAIL = "#诗雨会计"
STRIP_TAGS = ["#财会", "#备考日记", "#上班族备考", "#考前冲刺", "#备考焦虑"]

POSTS = [
    {
        "num": 232, "account": "沈小辉",
        "title": "书买了一摞真正翻开的没几本",
        "body": "备考资料买了一摞，真正翻开看过的没几本，钱花了心也虚。😅后来只留诗雨漫画笔记一本啃到底，反而学得进去。",
        "extra": ["#会计学习方法"],
    },
    {
        "num": 233, "account": "孙青5263",
        "title": "上班族每天能挤的时间太少",
        "body": "白天上班、晚上还有一堆事，真正能学习的时间少得可怜。😮‍💨后来靠诗雨漫画笔记利用零碎时间翻两页，也能一点点攒。",
        "extra": ["#上班族备考"],
    },
    {
        "num": 234, "account": "橙子学会计",
        "title": "同一个考点讲了三遍还是懵",
        "body": "同一个考点老师讲了好几遍，我还是听得云里雾里。😵后来用诗雨漫画笔记换个角度看，配着图一下就通了。",
        "extra": ["#中级会计实务"],
    },
    {
        "num": 235, "account": "新老婆",
        "title": "别人都开始刷题我还在看第一章",
        "body": "刷到别人都在刷题冲刺了，我还卡在第一章反复看。😣后来先用诗雨漫画笔记把高频考点快速过一遍，追进度没那么慌。",
        "extra": ["#中级会计备考"],
    },
    {
        "num": 236, "account": "孙玲小红书",
        "title": "看得懂视频课自己做就卡住",
        "body": "视频课跟着老师一步步都懂，自己独立做题就卡住动不了。😩后来靠诗雨漫画笔记自己复述考点，才发现哪里是假懂。",
        "extra": ["#会计学习方法"],
    },
    {
        "num": 237, "account": "爹爹",
        "title": "下班回家累到书都不想翻",
        "body": "上一天班回家整个人累趴，连翻书的力气都没有。😴后来靠诗雨漫画笔记，图多字少不费脑，躺着也能看两眼。",
        "extra": ["#上班族备考"],
    },
    {
        "num": 238, "account": "张菊香",
        "title": "笔记做得漂亮考点却没记住",
        "body": "笔记做得工工整整很好看，可考点还是没往脑子里去。😮‍💨后来发现诗雨漫画笔记本身就把重点画出来了，直接看更省事。",
        "extra": ["#会计学习方法"],
    },
    {
        "num": 239, "account": "孙文礼",
        "title": "一到大题就没思路只会小题",
        "body": "选择判断这些小题还行，一到综合大题就没思路、下不了笔。😣后来靠诗雨漫画笔记把答题框架理顺，大题也敢动手了。",
        "extra": ["#中级会计实务"],
    },
    {
        "num": 240, "account": "晴天",
        "title": "越临近考试越怕自己复习没到位",
        "body": "越接近考试越心慌，总怕自己有考点没复习到。😮‍💨后来对照诗雨漫画笔记的高频考点查漏补缺，心里才踏实。",
        "extra": ["#考前冲刺"],
    },
]

def check(title, body):
    tc = len(title)
    bc = len(body.replace("\n", "").replace(" ", ""))
    ok = "✅" if tc <= 20 and bc <= 80 else f"❌(标题{tc},正文{bc}字)"
    return ok, tc, bc

def gen_txts():
    for post in POSTS:
        title = post["title"]
        body = post["body"]
        tags = [P_BASE] + post["extra"] + [P_TAIL]
        topics = " ".join(tags)
        account = f"14芳芳财会|{post['account']}"
        ok, tc, bc = check(title, body)
        print(f"  {ok} 标题{tc}字 正文{bc}字 | {account}")
        keyword = title[:8]
        fname = f"{post['num']}_{keyword}_{post['account']}.txt"
        path = os.path.join(BASE, fname)
        content = f"标题：{title}\n\n正文：\n{body}\n\n话题：{topics}\n\n时间：\n\n账号：{account}\n"
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        post["_fname"] = fname

def gen_word():
    doc = Document()
    for i, post in enumerate(POSTS):
        body = post["body"]
        tags = [P_BASE] + post["extra"] + [P_TAIL]
        kept_tags = []
        for group in tags:
            for tag in group.split():
                if tag not in STRIP_TAGS and tag not in kept_tags:
                    kept_tags.append(tag)
        p1 = doc.add_paragraph()
        r1 = p1.add_run(body)
        r1.font.size = Pt(12)
        p2 = doc.add_paragraph()
        r2 = p2.add_run(" ".join(kept_tags))
        r2.font.size = Pt(12)
        if i != len(POSTS) - 1:
            p3 = doc.add_paragraph()
            run = p3.add_run()
            run.add_break(WD_BREAK.PAGE)
    path = os.path.join(BASE, DOCNAME)
    doc.save(path)
    print(f"  ✅ Word汇总已保存：{path}")

def make_zip():
    zip_path = os.path.join(BASE, ZIPNAME)
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for post in POSTS:
            zf.write(os.path.join(BASE, post["_fname"]), post["_fname"])
        zf.write(os.path.join(BASE, DOCNAME), DOCNAME)
    return zip_path

if __name__ == "__main__":
    print("=== 生成9篇引流txt ===")
    gen_txts()
    print("\n=== 生成Word汇总 ===")
    gen_word()
    zp = make_zip()
    print(f"\n  ✅ ZIP打包完成：{zp}")
