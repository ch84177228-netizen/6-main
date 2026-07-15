#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import os, zipfile
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_BREAK

BASE = "/home/user/6-main/01-6个主号内容创作/引流小号_发布文案"
P = "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考日记 #上班族备考 #考前冲刺 #诗雨会计"
DROP_TAGS = {"#财会", "#备考日记", "#上班族备考", "#考前冲刺"}

def save(path, content):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)

records = []

def write_txt(idx, title, body, account_name):
    tc = len(title)
    bc = len(body.replace("\n","").replace(" ",""))
    ok = "✅" if tc<=20 and bc<=80 else f"❌(标题{tc},正文{bc}字)"
    print(f"  {ok} 标题{tc}字 正文{bc}字 | {account_name}")
    kw = title[:8].replace("？","").replace("！","").replace("，","").replace("。","")
    fname = f"{idx}_{kw}_{account_name}.txt"
    content = f"标题：{title}\n\n正文：\n{body}\n\n话题：{P}\n\n时间：\n\n账号：14芳芳财会|{account_name}\n"
    save(os.path.join(BASE, fname), content)
    records.append((idx, body))

def gen_word(start, end):
    doc = Document()
    first = True
    for idx, body in records:
        if not first:
            p = doc.add_paragraph()
            p.add_run().add_break(WD_BREAK.PAGE)
        h = doc.add_paragraph()
        hr = h.add_run(f"第{idx}篇")
        hr.bold = True
        hr.font.size = Pt(14)
        for line in body.strip().split("\n"):
            p = doc.add_paragraph()
            r = p.add_run(line)
            r.font.size = Pt(12)
        tags = [t for t in P.split() if t not in DROP_TAGS]
        p = doc.add_paragraph()
        r = p.add_run(" ".join(tags))
        r.font.size = Pt(12)
        first = False
    out = os.path.join(BASE, f"引流小号_第{start}-{end}篇.docx")
    doc.save(out)
    print(f"  ✅ 汇总Word：{out}")

def make_zip(start, end):
    zip_path = os.path.join(BASE, f"引流小号_第{start}-{end}篇.zip")
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for i in range(start, end+1):
            for fname in os.listdir(BASE):
                if fname.startswith(f"{i}_") and fname.endswith(".txt"):
                    zf.write(os.path.join(BASE, fname), fname)
        wname = f"引流小号_第{start}-{end}篇.docx"
        if os.path.exists(os.path.join(BASE, wname)):
            zf.write(os.path.join(BASE, wname), wname)
    return zip_path

articles = [
    (142, "听课时全懂一做题就原形毕露",
     "听老师讲课频频点头，感觉全会了，一合上书做题就懵。😅原来听懂不等于会做。后来用诗雨漫画笔记自己复述一遍考点，才发现哪里是真懂哪里是假懂。",
     "沈小辉"),
    (143, "周末想补一整天结果睡到中午",
     "总想着周末一口气补一整天，结果一觉睡到中午，计划又黄了。😴后来改成每天固定翻几页诗雨漫画笔记，积少成多，比指望周末靠谱多了。",
     "孙青5263"),
    (144, "笔记记得很漂亮就是不会用",
     "笔记做得花花绿绿很好看，可一到考场脑子里啥也想不起来。😭花架子没用。换成诗雨漫画笔记重点突出，记的是真能用上的考点。",
     "橙子学会计"),
    (145, "总想等准备好了再开始结果一直没开始",
     "总觉得资料没买齐、状态没调好，等准备好再认真学，结果拖到现在还没真正开始。😣后来想通了，用诗雨漫画笔记先翻起来再说，开始了就不慌。",
     "新老婆"),
    (146, "做错的题讲过还是再错",
     "同一道题老师讲过、自己也标了，下次遇到照样错。😤说明根本没理解。配合诗雨漫画笔记把背后的原理搞懂，错题才真的不再错。",
     "孙玲小红书"),
    (147, "别人三个月上岸我半年还没动静",
     "看别人三个月就过了，我都准备半年了还没什么把握，心态有点崩。😮‍💨后来明白节奏不一样，用诗雨漫画笔记一点点啃，稳一点也能上岸。",
     "爹爹"),
    (148, "教材划满重点等于没划重点",
     "教材几乎每一行都划了线，结果等于没重点，复习时还是抓瞎。📚后来用诗雨漫画笔记，真正高频的考点一目了然，复习有了方向。",
     "张菊香"),
    (149, "背了一晚上第二天问就卡壳",
     "晚上背得滚瓜烂熟，第二天同事一问细节就卡壳，气得不行。😣机械背诵记不牢。用诗雨漫画笔记靠图理解记忆，第二天还能讲出来。",
     "孙文礼"),
    (150, "考前总幻想能押中题省点力",
     "总幻想押中考题能少看点书，结果押题没中基础又虚。😅还是踏实最稳。用诗雨漫画笔记把高频考点都过一遍，押不押中都不慌。",
     "晴天"),
]

if __name__ == "__main__":
    print("=== 生成引流小号第142-150篇 ===")
    for idx, title, body, account in articles:
        write_txt(idx, title, body, account)
    gen_word(142, 150)
    zp = make_zip(142, 150)
    print(f"\n  ✅ ZIP打包完成：{zp}")
