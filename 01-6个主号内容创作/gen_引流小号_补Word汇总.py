#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_BREAK

BASE = "/home/user/6-main/01-6个主号内容创作/引流小号_发布文案"

# Word汇总中需删除的话题标签（txt保持不变）
DROP_TAGS = {"#财会", "#备考日记", "#上班族备考", "#考前冲刺"}

# 全部批次：(起, 止)
BATCHES = [
    (22, 28), (29, 35), (36, 42), (43, 49), (50, 56),
    (57, 63), (64, 70), (71, 77), (78, 84),
    (85, 91), (92, 98), (99, 105), (106, 114), (115, 123),
]

def find_txt(idx):
    for fn in os.listdir(BASE):
        if fn.startswith(f"{idx}_") and fn.endswith(".txt"):
            return os.path.join(BASE, fn)
    return None

def parse_txt(content):
    """提取正文内容和话题标签（去掉标题/时间/账号字段，去掉'话题'二字）"""
    lines = content.split("\n")
    body_lines = []
    topic = ""
    section = None
    for line in lines:
        if line.startswith("标题："):
            section = "title"; continue
        if line.startswith("正文："):
            section = "body"; continue
        if line.startswith("话题："):
            raw = line[len("话题："):].strip()
            tags = [t for t in raw.split() if t not in DROP_TAGS]
            topic = " ".join(tags)
            section = "topic"; continue
        if line.startswith("时间：") or line.startswith("账号："):
            section = "other"; continue
        if section == "body":
            body_lines.append(line)
    # 去掉正文首尾空行
    while body_lines and body_lines[0].strip() == "":
        body_lines.pop(0)
    while body_lines and body_lines[-1].strip() == "":
        body_lines.pop()
    return "\n".join(body_lines), topic

def add_to_doc(doc, idx, body, topic, is_first):
    if not is_first:
        p = doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)
    h = doc.add_paragraph()
    hr = h.add_run(f"第{idx}篇")
    hr.bold = True
    hr.font.size = Pt(14)
    for line in body.split("\n"):
        p = doc.add_paragraph()
        r = p.add_run(line)
        r.font.size = Pt(12)
    if topic:
        p = doc.add_paragraph()
        r = p.add_run(topic)
        r.font.size = Pt(12)

def gen_batch(start, end):
    doc = Document()
    first = True
    count = 0
    for idx in range(start, end + 1):
        path = find_txt(idx)
        if not path:
            continue
        with open(path, "r", encoding="utf-8") as f:
            content = f.read()
        body, topic = parse_txt(content)
        add_to_doc(doc, idx, body, topic, first)
        first = False
        count += 1
    fname = f"引流小号_第{start}-{end}篇.docx"
    doc.save(os.path.join(BASE, fname))
    print(f"  ✅ {fname}（{count}篇）")

if __name__ == "__main__":
    print("=== 重新生成引流小号汇总Word（仅正文+话题标签）===")
    for s, e in BATCHES:
        gen_batch(s, e)
