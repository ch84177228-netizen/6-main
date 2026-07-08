#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import os, zipfile
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_BREAK

BASE = "/home/user/6-main/01-6个主号内容创作/引流小号_发布文案"
ZIPNAME = "编号20260708-01_引流账号文案_芳芳财会.zip"
DOCNAME = "引流小号_第187-195篇.docx"
P_BASE = "#会计 #中级会计 #中级会计备考 #注会cpa #财会"
P_TAIL = "#诗雨会计"
# Word 汇总去掉的标签（含本次明确要求的 #备考焦虑）
STRIP_TAGS = ["#财会", "#备考日记", "#上班族备考", "#考前冲刺", "#备考焦虑"]
# 本次禁止在任何 txt / Word 出现的标签
BANNED_TAGS = ["#备考焦虑"]

POSTS = [
    {
        "num": 187, "account": "沈小辉",
        "title": "通勤路上想背书结果眼睛太累",
        "body": "每天挤地铁想顺便背两页，结果晃得眼睛又酸又花，根本看不进去。😮‍💨后来换成诗雨漫画笔记，图多字少，通勤路上瞄两眼也不费眼。",
        "extra": ["#上班族备考"],
    },
    {
        "num": 188, "account": "孙青5263",
        "title": "不敢告诉家人自己在偷偷备考",
        "body": "报了名却没敢跟家里说，怕万一没考过被念叨。🤫只能自己默默学，压力全憋着。后来靠诗雨漫画笔记稳扎稳打，心里踏实了不少。",
        "extra": ["#零基础备考"],
    },
    {
        "num": 189, "account": "橙子学会计",
        "title": "直播课跟得上回放却直犯困",
        "body": "跟着直播课还挺精神，一到自己看回放就困得睁不开眼。😴后来搭配诗雨漫画笔记边看边对照，脑子跟着动起来，没那么容易走神。",
        "extra": ["#会计学习方法"],
    },
    {
        "num": 190, "account": "新老婆",
        "title": "思维导图画得漂亮却记不住",
        "body": "花一晚上画思维导图，图是好看，可考点还是记不牢。😅后来发现诗雨漫画笔记本身就把重点画出来了，直接看比自己画省事还好记。",
        "extra": ["#会计学习方法"],
    },
    {
        "num": 191, "account": "孙玲小红书",
        "title": "一遇到不会的就想从头再学",
        "body": "碰到一个不会的知识点，就忍不住想推倒从第一章重学，进度永远卡在开头。😩后来用诗雨漫画笔记哪里不会补哪里，才总算往前走了。",
        "extra": ["#中级会计实务"],
    },
    {
        "num": 192, "account": "爹爹",
        "title": "身边人都劝我别考了",
        "body": "同事都说这年纪考证没必要，听多了自己也动摇。😮‍💨但想想还是不甘心，跟着诗雨漫画笔记一点点啃，慢慢找回了点信心。",
        "extra": ["#上班族备考"],
    },
    {
        "num": 193, "account": "张菊香",
        "title": "同时报两科结果顾此失彼",
        "body": "贪心一次报了两科，结果这科学着那科忘着，两头都没学扎实。😣后来先用诗雨漫画笔记把各自高频考点理清，主次分明多了。",
        "extra": ["#中级会计备考"],
    },
    {
        "num": 194, "account": "孙文礼",
        "title": "学完一整遍感觉像没学过",
        "body": "教材从头翻到尾，合上书脑子一片空白，跟没学过一样。😵后来靠诗雨漫画笔记把框架串起来，再回看总算有印象了。",
        "extra": ["#会计学习方法"],
    },
    {
        "num": 195, "account": "晴天",
        "title": "越临近考试越想临时抱佛脚",
        "body": "眼看考试快到了，反而更想赌一把临时抱佛脚，心里其实没底。⏰后来老老实实用诗雨漫画笔记把高频考点过几遍，比乱抱佛脚安心多了。",
        "extra": ["#考前冲刺"],
    },
]

def check(title, body):
    tc = len(title)
    bc = len(body.replace("\n", "").replace(" ", ""))
    ok = "✅" if tc <= 20 and bc <= 80 else f"❌(标题{tc},正文{bc}字)"
    return ok, tc, bc

def gen_txts():
    for post in POSTS:
        title = post["title"]
        body = post["body"]
        tags = [P_BASE] + post["extra"] + [P_TAIL]
        # 确保禁用标签不出现
        assert all(bt not in " ".join(tags) for bt in BANNED_TAGS), f"禁用标签出现在 {post['num']}"
        topics = " ".join(tags)
        account = f"14芳芳财会|{post['account']}"
        ok, tc, bc = check(title, body)
        print(f"  {ok} 标题{tc}字 正文{bc}字 | {account}")
        keyword = title[:8]
        fname = f"{post['num']}_{keyword}_{post['account']}.txt"
        path = os.path.join(BASE, fname)
        content = f"标题：{title}\n\n正文：\n{body}\n\n话题：{topics}\n\n时间：\n\n账号：{account}\n"
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        post["_fname"] = fname

def gen_word():
    doc = Document()
    for i, post in enumerate(POSTS):
        body = post["body"]
        tags = [P_BASE] + post["extra"] + [P_TAIL]
        kept_tags = []
        for group in tags:
            for tag in group.split():
                if tag not in STRIP_TAGS and tag not in kept_tags:
                    kept_tags.append(tag)
        p1 = doc.add_paragraph()
        r1 = p1.add_run(body)
        r1.font.size = Pt(12)
        p2 = doc.add_paragraph()
        r2 = p2.add_run(" ".join(kept_tags))
        r2.font.size = Pt(12)
        if i != len(POSTS) - 1:
            p3 = doc.add_paragraph()
            run = p3.add_run()
            run.add_break(WD_BREAK.PAGE)
    path = os.path.join(BASE, DOCNAME)
    doc.save(path)
    print(f"  ✅ Word汇总已保存：{path}")

def make_zip():
    zip_path = os.path.join(BASE, ZIPNAME)
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for post in POSTS:
            zf.write(os.path.join(BASE, post["_fname"]), post["_fname"])
        zf.write(os.path.join(BASE, DOCNAME), DOCNAME)
    return zip_path

if __name__ == "__main__":
    print("=== 生成9篇引流txt ===")
    gen_txts()
    print("\n=== 生成Word汇总 ===")
    gen_word()
    zp = make_zip()
    print(f"\n  ✅ ZIP打包完成：{zp}")
