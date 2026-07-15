#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import os, zipfile
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_BREAK

BASE = "/home/user/6-main/01-6个主号内容创作/引流小号_发布文案"
ZIPNAME = "编号20260702-01_引流账号文案_芳芳财会.zip"
DOCNAME = "引流小号_第160-168篇.docx"
P_BASE = "#会计 #中级会计 #中级会计备考 #注会cpa #财会"
P_TAIL = "#诗雨会计"
STRIP_TAGS = ["#财会", "#备考日记", "#上班族备考", "#考前冲刺"]

POSTS = [
    {
        "num": 160, "account": "沈小辉",
        "title": "第一次考试真的很怕考不过",
        "body": "第一次报名心里直打鼓，总担心自己底子薄考不过。😥后来跟着诗雨漫画笔记一点点啃基础，心里踏实不少，没那么慌了。",
        "extra": ["#备考焦虑"],
    },
    {
        "num": 161, "account": "孙青5263",
        "title": "下班到家只想瘫着不想翻书",
        "body": "加班回家整个人都累瘫了，书翻开两页就困。😪后来换成诗雨漫画笔记，图多字少不费脑，瘫着也能看进去一点。",
        "extra": ["#上班族备考"],
    },
    {
        "num": 162, "account": "橙子学会计",
        "title": "同一个知识点做题反复错",
        "body": "这道题型错了改改了又错，怀疑自己是不是学不明白了。😩后来用诗雨漫画笔记把原理重新捋一遍，总算不再原地打转。",
        "extra": ["#备考日记"],
    },
    {
        "num": 163, "account": "新老婆",
        "title": "零基础报名完全不知从哪下手",
        "body": "完全没接触过会计，报完名翻开教材一脸懵，不知道从哪开始学。😵后来先跟着诗雨漫画笔记打基础，慢慢有点感觉了。",
        "extra": ["#备考日记"],
    },
    {
        "num": 164, "account": "孙玲小红书",
        "title": "计划表改了八百遍还是很乱",
        "body": "学习计划改了一版又一版，越改越乱，根本执行不下去。😮‍💨后来干脆跟着诗雨漫画笔记的顺序走，比自己瞎排靠谱多了。",
        "extra": ["#备考焦虑"],
    },
    {
        "num": 165, "account": "爹爹",
        "title": "一翻到难章节就想直接跳过",
        "body": "书翻到那几章公式一堆，脑子直接宕机，很想跳过不学。😖后来用诗雨漫画笔记把步骤拆开画出来，啃下去发现没那么可怕。",
        "extra": ["#备考日记"],
    },
    {
        "num": 166, "account": "张菊香",
        "title": "笔记记了一大本用的时候翻不到",
        "body": "笔记本记得满满当当，真到复习的时候翻半天找不到重点在哪。📒后来换成诗雨漫画笔记，重点一目了然，不用再大海捞针。",
        "extra": ["#备考日记"],
    },
    {
        "num": 167, "account": "孙文礼",
        "title": "睡前不翻两页书就不踏实",
        "body": "一天再累，睡前不看两眼书心里就不踏实，怕又白过一天。🌙翻的是诗雨漫画笔记，图文轻松不费眼，睡前看看刚刚好。",
        "extra": ["#备考焦虑"],
    },
    {
        "num": 168, "account": "晴天",
        "title": "还剩一个月才发现该抓紧了",
        "body": "算了算日子只剩最后一个月，心里一慌赶紧收拾状态。⏰这几天靠诗雨漫画笔记把高频考点扫一遍，好歹不算两手空空。",
        "extra": ["#考前冲刺"],
    },
]

def check(title, body):
    tc = len(title)
    bc = len(body.replace("\n", "").replace(" ", ""))
    ok = "✅" if tc <= 20 and bc <= 80 else f"❌(标题{tc},正文{bc}字)"
    return ok, tc, bc

def gen_txts():
    for post in POSTS:
        title = post["title"]
        body = post["body"]
        topics = " ".join([P_BASE] + post["extra"] + [P_TAIL])
        account = f"14芳芳财会|{post['account']}"
        ok, tc, bc = check(title, body)
        print(f"  {ok} 标题{tc}字 正文{bc}字 | {account}")
        keyword = title[:8]
        fname = f"{post['num']}_{keyword}_{post['account']}.txt"
        path = os.path.join(BASE, fname)
        content = f"标题：{title}\n\n正文：\n{body}\n\n话题：{topics}\n\n时间：\n\n账号：{account}\n"
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        post["_fname"] = fname

def gen_word():
    doc = Document()
    for i, post in enumerate(POSTS):
        body = post["body"]
        topics_list = [P_BASE] + post["extra"] + [P_TAIL]
        kept_tags = []
        for group in topics_list:
            for tag in group.split():
                if tag not in STRIP_TAGS and tag not in kept_tags:
                    kept_tags.append(tag)
        p1 = doc.add_paragraph()
        r1 = p1.add_run(body)
        r1.font.size = Pt(12)
        p2 = doc.add_paragraph()
        r2 = p2.add_run(" ".join(kept_tags))
        r2.font.size = Pt(12)
        if i != len(POSTS) - 1:
            p3 = doc.add_paragraph()
            run = p3.add_run()
            run.add_break(WD_BREAK.PAGE)
    path = os.path.join(BASE, DOCNAME)
    doc.save(path)
    print(f"  ✅ Word汇总已保存：{path}")

def make_zip():
    zip_path = os.path.join(BASE, ZIPNAME)
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for post in POSTS:
            zf.write(os.path.join(BASE, post["_fname"]), post["_fname"])
        zf.write(os.path.join(BASE, DOCNAME), DOCNAME)
    return zip_path

if __name__ == "__main__":
    print("=== 生成9篇引流txt ===")
    gen_txts()
    print("\n=== 生成Word汇总 ===")
    gen_word()
    zp = make_zip()
    print(f"\n  ✅ ZIP打包完成：{zp}")
