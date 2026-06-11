"""
引流小号 第22-28篇生成脚本
已有主题（避免重复）：
1-7: 下班后学不动/资料买太多/教材想睡/没出门/没时间/全记错了/考前两月救吗
8-14: 不知从哪下手/做题错误率高/背了就忘/别人进度快/孩子太闹/碎片时间/备考焦虑失眠
15-21: 目录没背完87天/同事问进度不敢说/计划表排5遍没完成/看课懂了做题全错/初级3年才备考中级/午休偷学被发现/刷了300题分数上不去
"""

import os
import zipfile
from docx import Document

BASE = "/home/user/6-main/01-6个主号内容创作/引流小号_发布文案"

articles = [
    {
        "no": 22,
        "account": "沈小辉",
        "keyword": "报名费交了不去考太亏了",
        "title": "报名费交了不去考还是太亏了",
        "body": (
            "当初报名就是随便试试😅\n"
            "交完钱才算了算，三科好几百\n"
            "\n"
            "这几百块反而成了最大动力\n"
            "每次想摆烂就算一遍报名费💀\n"
            "\n"
            "最近拿诗雨漫画笔记看\n"
            "图多，起码看得下去不白交钱"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考日记 #上班族备考 #诗雨会计",
    },
    {
        "no": 23,
        "account": "孙青5263",
        "keyword": "工作忙只剩睡前三十分钟",
        "title": "工作太忙只剩睡前那三十分钟",
        "body": (
            "下班回家已经八点多了😮‍💨\n"
            "吃饭洗澡，真正能坐下来快十点\n"
            "\n"
            "给自己留了睡前30分钟学会计\n"
            "翻了三页眼皮开始打架💤\n"
            "\n"
            "后来改看诗雨漫画笔记\n"
            "图多字少，30分钟翻完一节"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #上班族备考 #备考日记 #诗雨会计",
    },
    {
        "no": 24,
        "account": "橙子学会计",
        "keyword": "手写两本笔记发现没用上",
        "title": "手写了满满两本笔记发现没用上",
        "body": (
            "备考第一件事就是抄笔记✍️\n"
            "写了两本，字越写越好看\n"
            "\n"
            "做题才发现完全对不上号\n"
            "好像只是把教材抄了一遍😑\n"
            "\n"
            "后来换成诗雨漫画笔记\n"
            "本来就有图，不用自己再抄了"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考方法 #备考日记 #诗雨会计",
    },
    {
        "no": 25,
        "account": "新老婆",
        "keyword": "打开教材5分钟就刷手机",
        "title": "打开教材看了5分钟就开始刷手机",
        "body": (
            "下定决心学会计✊\n"
            "打开教材，好，开始学\n"
            "\n"
            "5分钟后——刷手机😅\n"
            "10分钟后——还在刷\n"
            "半小时后——还是在刷\n"
            "\n"
            "后来改用诗雨漫画笔记\n"
            "漫画版，多看两页都行"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考焦虑 #备考方法 #诗雨会计",
    },
    {
        "no": 26,
        "account": "孙玲小红书",
        "keyword": "上班路上背题坐过站了",
        "title": "上班路上背题结果坐过站了",
        "body": (
            "地铁上刷题刷到入神🚇\n"
            "抬头一看：终点站\n"
            "\n"
            "多坐了六站，迟到了\n"
            "同事问怎么了不敢说😶\n"
            "\n"
            "后来换成诗雨漫画笔记\n"
            "坐过站了就当多看一节"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #上班族备考 #备考日记 #诗雨会计",
    },
    {
        "no": 27,
        "account": "爹爹",
        "keyword": "报了辅导班比自学还慌",
        "title": "报了辅导班发现比自学还慌",
        "body": (
            "花钱报了辅导班💸\n"
            "以为付了钱就稳了\n"
            "\n"
            "第一节课老师讲了80页\n"
            "第二节课直接跳第三章😵\n"
            "\n"
            "跟不上，自学也没时间\n"
            "后来拿诗雨漫画笔记倒着翻，好一点🥲"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考方法 #备考焦虑 #诗雨会计",
    },
    {
        "no": 28,
        "account": "张菊香",
        "keyword": "越临近考试越不想翻书",
        "title": "越临近考试越不想翻书是啥心态",
        "body": (
            "距考试还有86天\n"
            "突然不想打开教材了😶\n"
            "\n"
            "反而去刷视频、刷剧、刷一切\n"
            "就是不想看会计\n"
            "\n"
            "后来查了下好像叫\"考前逃避\"😅\n"
            "换成诗雨漫画笔记随便翻翻\n"
            "好歹动起来了"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #考前冲刺 #备考焦虑 #诗雨会计",
    },
]

def make_txt(a):
    return (
        f"标题：{a['title']}\n"
        f"\n"
        f"正文：\n"
        f"{a['body']}\n"
        f"\n"
        f"话题：{a['tags']}\n"
        f"\n"
        f"时间：\n"
        f"\n"
        f"账号：14芳芳财会|{a['account']}\n"
    )

txt_paths = []
for a in articles:
    fname = f"{a['no']}_{a['keyword']}_{a['account']}.txt"
    fpath = os.path.join(BASE, fname)
    with open(fpath, "w", encoding="utf-8") as f:
        f.write(make_txt(a))
    txt_paths.append(fpath)
    print(f"✅ {fname}")

# Word document
doc = Document()
doc.add_heading("引流小号文案 第22-28篇", 0)
for a in articles:
    doc.add_heading(f"{a['no']}. 账号：14芳芳财会|{a['account']}", 1)
    p = doc.add_paragraph()
    p.add_run("标题：").bold = True
    p.add_run(a["title"])
    doc.add_paragraph()
    doc.add_paragraph("正文：")
    doc.add_paragraph(a["body"])
    doc.add_paragraph()
    p2 = doc.add_paragraph()
    p2.add_run("话题：").bold = False
    p2.add_run(a["tags"])
    doc.add_paragraph(f"账号：14芳芳财会|{a['account']}")
    doc.add_paragraph("—" * 30)

word_path = os.path.join(BASE, "引流小号_第22-28篇.docx")
doc.save(word_path)
print(f"✅ Word: {word_path}")

# ZIP
zip_path = os.path.join(BASE, "引流小号_第22-28篇_7篇文案.zip")
with zipfile.ZipFile(zip_path, "w") as zf:
    for fp in txt_paths:
        zf.write(fp, os.path.basename(fp))
print(f"✅ ZIP: {zip_path}")
print("全部完成！")
