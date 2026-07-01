#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import os, zipfile
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_BREAK

BASE = "/home/user/6-main/01-6个主号内容创作/引流小号_发布文案"
ZIPNAME = "编号20260701-11_引流账号文案_芳芳财会.zip"
DOCNAME = "引流小号_第151-159篇.docx"
P_BASE = "#会计 #中级会计 #中级会计备考 #注会cpa #财会"
P_TAIL = "#诗雨会计"
STRIP_TAGS = ["#财会", "#备考日记", "#上班族备考", "#考前冲刺"]

POSTS = [
    {
        "num": 151, "account": "沈小辉",
        "title": "手机刷题刷着刷着刷起短视频",
        "body": "想着刷题提神，结果越刷越忍不住点开短视频，效率直接归零。😩后来把手机丢一边，用诗雨漫画笔记刷考点，眼睛没地方跑神。",
        "extra": ["#备考日记", "#上班族备考"],
    },
    {
        "num": 152, "account": "孙青5263",
        "title": "特意请假在家结果一天没学进去",
        "body": "特意请了一天假想在家猛学，结果东摸西摸，一天下来没学进去几页。😮‍💨后来靠诗雨漫画笔记定量看，不请假也能稳步推进。",
        "extra": ["#备考焦虑"],
    },
    {
        "num": 153, "account": "橙子学会计",
        "title": "两个相似概念做题总是选反",
        "body": "两个听起来很像的概念，一到选择题就选反，反复错好几次。😵后来用诗雨漫画笔记把区别画出来对比看，总算分清楚了。",
        "extra": ["#备考日记"],
    },
    {
        "num": 154, "account": "新老婆",
        "title": "深夜emo怀疑自己学不会",
        "body": "白天还挺有干劲，一到深夜就emo，怀疑自己是不是学不会了。😔后来把诗雨漫画笔记翻两页，看懂一个考点，心里踏实点。",
        "extra": ["#备考焦虑", "#上班族备考"],
    },
    {
        "num": 155, "account": "孙玲小红书",
        "title": "计划表做得完美一天没跟上",
        "body": "计划表排得整整齐齐，看着很有仪式感，结果第一天就没跟上进度。😅后来换成诗雨漫画笔记按考点走，不硬套计划反而更稳。",
        "extra": ["#备考日记"],
    },
    {
        "num": 156, "account": "爹爹",
        "title": "上班间隙偷学两页也算学习",
        "body": "上班间隙偷偷翻两页复习，被同事打趣像做贼一样。😂能学一点是一点。诗雨漫画笔记图多字少，碎片时间翻两页也能记住点东西。",
        "extra": ["#上班族备考"],
    },
    {
        "num": 157, "account": "张菊香",
        "title": "报了两门课时间根本不够分",
        "body": "同时报了两门课，时间怎么分都不够用，越学越慌。😣后来先用诗雨漫画笔记把高频考点过一遍，主次分清楚才没那么乱。",
        "extra": ["#备考焦虑"],
    },
    {
        "num": 158, "account": "孙文礼",
        "title": "同一道题反复出现还是不会",
        "body": "这道题型都出现好几次了，每次遇到还是卡壳，挫败感很强。😤后来靠诗雨漫画笔记把原理理清楚，再遇到才敢下笔。",
        "extra": ["#备考日记"],
    },
    {
        "num": 159, "account": "晴天",
        "title": "考前失眠越紧张越睡不着",
        "body": "越接近考试越紧张，晚上翻来覆去睡不着，白天没精神。😴后来睡前翻两页诗雨漫画笔记，看着看着反而放松了不少。",
        "extra": ["#考前冲刺"],
    },
]

def check(title, body):
    tc = len(title)
    bc = len(body.replace("\n", "").replace(" ", ""))
    ok = "✅" if tc <= 20 and bc <= 80 else f"❌(标题{tc},正文{bc}字)"
    return ok, tc, bc

def gen_txts():
    for post in POSTS:
        title = post["title"]
        body = post["body"]
        topics = " ".join([P_BASE] + post["extra"] + [P_TAIL])
        account = f"14芳芳财会|{post['account']}"
        ok, tc, bc = check(title, body)
        print(f"  {ok} 标题{tc}字 正文{bc}字 | {account}")
        keyword = title[:8]
        fname = f"{post['num']}_{keyword}_{post['account']}.txt"
        path = os.path.join(BASE, fname)
        content = f"标题：{title}\n\n正文：\n{body}\n\n话题：{topics}\n\n时间：\n\n账号：{account}\n"
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        post["_fname"] = fname

def gen_word():
    doc = Document()
    for i, post in enumerate(POSTS):
        body = post["body"]
        topics_list = [P_BASE] + post["extra"] + [P_TAIL]
        kept_tags = []
        for group in topics_list:
            for tag in group.split():
                if tag not in STRIP_TAGS and tag not in kept_tags:
                    kept_tags.append(tag)
        p1 = doc.add_paragraph()
        r1 = p1.add_run(body)
        r1.font.size = Pt(12)
        p2 = doc.add_paragraph()
        r2 = p2.add_run(" ".join(kept_tags))
        r2.font.size = Pt(12)
        if i != len(POSTS) - 1:
            p3 = doc.add_paragraph()
            run = p3.add_run()
            run.add_break(WD_BREAK.PAGE)
    path = os.path.join(BASE, DOCNAME)
    doc.save(path)
    print(f"  ✅ Word汇总已保存：{path}")

def make_zip():
    zip_path = os.path.join(BASE, ZIPNAME)
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for post in POSTS:
            zf.write(os.path.join(BASE, post["_fname"]), post["_fname"])
        zf.write(os.path.join(BASE, DOCNAME), DOCNAME)
    return zip_path

if __name__ == "__main__":
    print("=== 生成9篇引流txt ===")
    gen_txts()
    print("\n=== 生成Word汇总 ===")
    gen_word()
    zp = make_zip()
    print(f"\n  ✅ ZIP打包完成：{zp}")
