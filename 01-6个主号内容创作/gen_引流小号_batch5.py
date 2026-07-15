"""
引流小号 第36-42篇生成脚本
已有主题（避免重复）：
1-7: 下班学不动/资料太多/教材想睡/没出门/没时间复习/全记错了/考前两月救吗
8-14: 不知从哪下手/错误率高崩溃/背了就忘/别人进度快/孩子太闹/碎片时间/焦虑失眠
15-21: 目录没背完87天/不敢说进度/计划表5遍没完/看课懂了做题全错/初级3年后备考/午休被发现/刷300题分数上不去
22-28: 报名费不考亏/睡前30分钟/手写笔记没用/开教材刷手机/背题坐过站/辅导班更慌/越近越不想翻书
29-35: 考前全忘了/公式背了忘/家人催问/看别人打卡慌/出差断学/看了个寂寞/模拟题错一半
"""

import os
import zipfile
from docx import Document

BASE = "/home/user/6-main/01-6个主号内容创作/引流小号_发布文案"

articles = [
    {
        "no": 36,
        "account": "沈小辉",
        "keyword": "备考到一半突然换教材了",
        "title": "备考到一半突然决定换教材",
        "body": (
            "备考三个月了\n"
            "朋友说她那套教材更好用👀\n"
            "\n"
            "于是我花了一周时间换教材\n"
            "重新从第一章翻起😑\n"
            "\n"
            "后来没再换了\n"
            "拿诗雨漫画笔记配着原来那套走完的"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考方法 #备考日记 #诗雨会计",
    },
    {
        "no": 37,
        "account": "孙青5263",
        "keyword": "同一科报了两家机构全没跟完",
        "title": "同一门课报了两家机构两家都没跟完",
        "body": (
            "当初觉得一家不够稳妥💸\n"
            "同一门课报了两家辅导班\n"
            "\n"
            "结果两家的课表撞了\n"
            "A家没跟完，B家也没跟完😮‍💨\n"
            "\n"
            "最后跟着诗雨漫画笔记自学\n"
            "省钱，还跟得完"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考方法 #备考焦虑 #诗雨会计",
    },
    {
        "no": 38,
        "account": "橙子学会计",
        "keyword": "手机备忘录存了100条考点没打开",
        "title": "手机备忘录存了100条考点从没打开",
        "body": (
            "备考最勤快的事情就是存资料📱\n"
            "备忘录、收藏夹、云盘……\n"
            "\n"
            "临考前翻了翻：100条笔记，0条看过\n"
            "资料存了个寂寞😶\n"
            "\n"
            "后来直接拿诗雨漫画笔记翻\n"
            "实体的，翻了就是翻了"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考方法 #备考日记 #诗雨会计",
    },
    {
        "no": 39,
        "account": "新老婆",
        "keyword": "模拟考偷看答案自我感觉不错",
        "title": "模拟考偷看着答案感觉学得不错",
        "body": (
            "刷模拟题的时候习惯边看答案边做😅\n"
            "感觉做起来很顺\n"
            "\n"
            "直到真正闭卷测了一次\n"
            "第一题——空白😶\n"
            "\n"
            "换成诗雨漫画笔记重新理一遍框架\n"
            "闭卷才算真的会"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #刷题 #备考方法 #诗雨会计",
    },
    {
        "no": 40,
        "account": "孙玲小红书",
        "keyword": "考前一周发现有一章完全没看",
        "title": "考前一周发现整整一章完全没看过",
        "body": (
            "考前一周翻目录做检查✅\n"
            "突然看到第七章\n"
            "\n"
            "我：这章讲啥来着？\n"
            "打开教材：全是生面孔😱\n"
            "\n"
            "赶紧拿诗雨漫画笔记扫了一遍\n"
            "考前一周，图解版还能抢救"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #考前冲刺 #备考焦虑 #诗雨会计",
    },
    {
        "no": 41,
        "account": "爹爹",
        "keyword": "每次学习前要整理桌面整理一小时",
        "title": "每次准备学习都要先整理一小时桌面",
        "body": (
            "计划8点开始学\n"
            "7点55分：先整理下桌面🧹\n"
            "\n"
            "8点30分：桌面收拾好了\n"
            "顺手擦了擦窗台\n"
            "9点：开始正式整理椅子……\n"
            "\n"
            "后来把诗雨漫画笔记直接放床头\n"
            "不用桌面，直接翻"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考日记 #备考焦虑 #诗雨会计",
    },
    {
        "no": 42,
        "account": "张菊香",
        "keyword": "差三分没过的那一年",
        "title": "差三分没过的那年，我到底差在哪",
        "body": (
            "去年差三分\n"
            "成绩出来那天盯着屏幕看了很久😔\n"
            "\n"
            "复盘了很久：不是不努力\n"
            "是学的方法跟考的方式对不上\n"
            "\n"
            "今年换了诗雨漫画笔记重新备考\n"
            "按考点学，不按章节学了"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考方法 #考前冲刺 #诗雨会计",
    },
]

def make_txt(a):
    return (
        f"标题：{a['title']}\n"
        f"\n"
        f"正文：\n"
        f"{a['body']}\n"
        f"\n"
        f"话题：{a['tags']}\n"
        f"\n"
        f"时间：\n"
        f"\n"
        f"账号：14芳芳财会|{a['account']}\n"
    )

txt_paths = []
for a in articles:
    fname = f"{a['no']}_{a['keyword']}_{a['account']}.txt"
    fpath = os.path.join(BASE, fname)
    with open(fpath, "w", encoding="utf-8") as f:
        f.write(make_txt(a))
    txt_paths.append(fpath)
    print(f"✅ {fname}")

# Word document
doc = Document()
doc.add_heading("引流小号文案 第36-42篇", 0)
for a in articles:
    doc.add_heading(f"{a['no']}. 账号：14芳芳财会|{a['account']}", 1)
    p = doc.add_paragraph()
    p.add_run("标题：").bold = True
    p.add_run(a["title"])
    doc.add_paragraph()
    doc.add_paragraph("正文：")
    doc.add_paragraph(a["body"])
    doc.add_paragraph()
    p2 = doc.add_paragraph()
    p2.add_run("话题：")
    p2.add_run(a["tags"])
    doc.add_paragraph(f"账号：14芳芳财会|{a['account']}")
    doc.add_paragraph("—" * 30)

word_path = os.path.join(BASE, "引流小号_第36-42篇.docx")
doc.save(word_path)
print(f"✅ Word: {word_path}")

# ZIP（txt only）
zip_path = os.path.join(BASE, "引流小号_第36-42篇_7篇文案.zip")
with zipfile.ZipFile(zip_path, "w") as zf:
    for fp in txt_paths:
        zf.write(fp, os.path.basename(fp))
print(f"✅ ZIP: {zip_path}")
print("全部完成！")
