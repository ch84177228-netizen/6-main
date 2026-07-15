#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import os, zipfile
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_BREAK

BASE = "/home/user/6-main/01-6个主号内容创作/引流小号_发布文案"
ZIPNAME = "编号20260704-01_引流账号文案_芳芳财会.zip"
DOCNAME = "引流小号_第169-177篇.docx"
P_BASE = "#会计 #中级会计 #中级会计备考 #注会cpa #财会"
P_TAIL = "#诗雨会计"
STRIP_TAGS = ["#财会", "#备考日记", "#上班族备考", "#考前冲刺"]

POSTS = [
    {
        "num": 169, "account": "沈小辉",
        "title": "加班到很晚回家还想学但没劲",
        "body": "加班到很晚，回家满脑子只想躺平，可一想到进度又不甘心。😮‍💨后来靠诗雨漫画笔记翻两页，不费脑子也能学进去一点。",
        "extra": ["#上班族备考"],
    },
    {
        "num": 170, "account": "孙青5263",
        "title": "报了班总能找借口拖延",
        "body": "明明报了班，还是能给自己找一百个理由往后拖。😅后来把诗雨漫画笔记放手边，翻开就学，少了纠结的时间。",
        "extra": ["#备考日记"],
    },
    {
        "num": 171, "account": "橙子学会计",
        "title": "刷题正确率忽高忽低很崩溃",
        "body": "这次刷题全对，下次同类题又错一堆，心态很崩。😵后来用诗雨漫画笔记把原理吃透，正确率才慢慢稳住。",
        "extra": ["#备考焦虑"],
    },
    {
        "num": 172, "account": "新老婆",
        "title": "家里太吵根本静不下心学习",
        "body": "家里孩子闹、电视响，想安静学一会儿都难。😩后来靠诗雨漫画笔记，图多字少，吵一点也能瞄两眼记住点东西。",
        "extra": ["#备考日记"],
    },
    {
        "num": 173, "account": "孙玲小红书",
        "title": "总想换新教材结果越换越乱",
        "body": "总觉得换本新教材能学得更好，结果换来换去反而更乱。😮‍💨后来定下诗雨漫画笔记一本啃到底，思路清楚多了。",
        "extra": ["#备考日记"],
    },
    {
        "num": 174, "account": "爹爹",
        "title": "白天上班脑子转不动只能靠晚上",
        "body": "白天上班脑子完全转不动，只能指望晚上挤出点时间。😴晚上靠诗雨漫画笔记效率高一些，好歹没白熬。",
        "extra": ["#上班族备考"],
    },
    {
        "num": 175, "account": "张菊香",
        "title": "刷到别人进度贴突然emo了",
        "body": "刷到别人晒的复习进度，瞬间觉得自己落后一大截。😔后来不比较了，跟着诗雨漫画笔记按自己节奏走，稳一点。",
        "extra": ["#备考焦虑"],
    },
    {
        "num": 176, "account": "孙文礼",
        "title": "审题不仔细总丢冤枉分",
        "body": "明明会做，就因为审题不仔细丢了冤枉分，气死自己。😤后来用诗雨漫画笔记多看几遍易错点，细节抓得更准了。",
        "extra": ["#备考日记"],
    },
    {
        "num": 177, "account": "晴天",
        "title": "考前紧张手心出汗写字发抖",
        "body": "一想到要考试就紧张，手心冒汗写字都在抖。😰后来靠诗雨漫画笔记把高频考点多过几遍，心里踏实点没那么慌。",
        "extra": ["#考前冲刺"],
    },
]

def check(title, body):
    tc = len(title)
    bc = len(body.replace("\n", "").replace(" ", ""))
    ok = "✅" if tc <= 20 and bc <= 80 else f"❌(标题{tc},正文{bc}字)"
    return ok, tc, bc

def gen_txts():
    for post in POSTS:
        title = post["title"]
        body = post["body"]
        topics = " ".join([P_BASE] + post["extra"] + [P_TAIL])
        account = f"14芳芳财会|{post['account']}"
        ok, tc, bc = check(title, body)
        print(f"  {ok} 标题{tc}字 正文{bc}字 | {account}")
        keyword = title[:8]
        fname = f"{post['num']}_{keyword}_{post['account']}.txt"
        path = os.path.join(BASE, fname)
        content = f"标题：{title}\n\n正文：\n{body}\n\n话题：{topics}\n\n时间：\n\n账号：{account}\n"
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        post["_fname"] = fname

def gen_word():
    doc = Document()
    for i, post in enumerate(POSTS):
        body = post["body"]
        topics_list = [P_BASE] + post["extra"] + [P_TAIL]
        kept_tags = []
        for group in topics_list:
            for tag in group.split():
                if tag not in STRIP_TAGS and tag not in kept_tags:
                    kept_tags.append(tag)
        p1 = doc.add_paragraph()
        r1 = p1.add_run(body)
        r1.font.size = Pt(12)
        p2 = doc.add_paragraph()
        r2 = p2.add_run(" ".join(kept_tags))
        r2.font.size = Pt(12)
        if i != len(POSTS) - 1:
            p3 = doc.add_paragraph()
            run = p3.add_run()
            run.add_break(WD_BREAK.PAGE)
    path = os.path.join(BASE, DOCNAME)
    doc.save(path)
    print(f"  ✅ Word汇总已保存：{path}")

def make_zip():
    zip_path = os.path.join(BASE, ZIPNAME)
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for post in POSTS:
            zf.write(os.path.join(BASE, post["_fname"]), post["_fname"])
        zf.write(os.path.join(BASE, DOCNAME), DOCNAME)
    return zip_path

if __name__ == "__main__":
    print("=== 生成9篇引流txt ===")
    gen_txts()
    print("\n=== 生成Word汇总 ===")
    gen_word()
    zp = make_zip()
    print(f"\n  ✅ ZIP打包完成：{zp}")
