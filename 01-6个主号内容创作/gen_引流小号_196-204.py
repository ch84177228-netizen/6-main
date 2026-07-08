#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import os, zipfile
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_BREAK

BASE = "/home/user/6-main/01-6个主号内容创作/引流小号_发布文案"
ZIPNAME = "编号20260708-02_引流账号文案_芳芳财会.zip"
DOCNAME = "引流小号_第196-204篇.docx"
P_BASE = "#会计 #中级会计 #中级会计备考 #注会cpa #财会"
P_TAIL = "#诗雨会计"
STRIP_TAGS = ["#财会", "#备考日记", "#上班族备考", "#考前冲刺", "#备考焦虑"]

POSTS = [
    {
        "num": 196, "account": "沈小辉",
        "title": "买了一堆资料结果全新落灰",
        "body": "冲动买了一堆备考资料，翻了没几页就堆在角落落灰。😅后来只留诗雨漫画笔记一本，翻起来不累也不会再囤新的了。",
        "extra": ["#会计学习方法"],
    },
    {
        "num": 197, "account": "孙青5263",
        "title": "看着目录厚就先放弃了一半",
        "body": "打开教材看到目录那么厚，还没开始学就先泄气了一半。😮‍💨后来跟着诗雨漫画笔记一小块一小块过，反而没那么怕了。",
        "extra": ["#零基础备考"],
    },
    {
        "num": 198, "account": "橙子学会计",
        "title": "分录背了又忘忘了又背",
        "body": "会计分录背了忘、忘了再背，来回折腾好几遍还是记不牢。😩后来靠诗雨漫画笔记把分录的逻辑理顺，理解着记就没那么容易忘了。",
        "extra": ["#中级会计实务"],
    },
    {
        "num": 199, "account": "新老婆",
        "title": "一学习就开始困上头",
        "body": "白天精神好好的，一坐下来看书就开始犯困、上头。😴后来发现诗雨漫画笔记图多字少，看着轻松，反而能多撑一会儿。",
        "extra": ["#会计学习方法"],
    },
    {
        "num": 200, "account": "孙玲小红书",
        "title": "总觉得时间还早就一直拖",
        "body": "每次都想着离考试还早，明天再学也来得及，结果一天天就拖过去了。😮‍💨后来每天翻两页诗雨漫画笔记，积少成多才不慌。",
        "extra": ["#中级会计备考"],
    },
    {
        "num": 201, "account": "爹爹",
        "title": "下班太累根本坐不住看书",
        "body": "上一天班回家整个人瘫着，让我端坐看书实在坐不住。😮‍💨后来靠诗雨漫画笔记，躺着翻两页也能记住点东西，压力小多了。",
        "extra": ["#上班族备考"],
    },
    {
        "num": 202, "account": "张菊香",
        "title": "看懂了例题一换数字就不会",
        "body": "老师讲的例题都听懂了，自己一换个数字就又不会做了。😣后来用诗雨漫画笔记把解题步骤拆开记，换汤不换药也能应付了。",
        "extra": ["#中级会计实务"],
    },
    {
        "num": 203, "account": "孙文礼",
        "title": "重点划满一整本等于没划",
        "body": "教材几乎每页都划满了重点，结果一到复习还是抓不住关键。📚后来换成诗雨漫画笔记，真正高频的考点一目了然，复习有方向了。",
        "extra": ["#会计学习方法"],
    },
    {
        "num": 204, "account": "晴天",
        "title": "越到最后越怕之前白学了",
        "body": "越接近考试越心虚，怕前面学的都白费了。😮‍💨后来靠诗雨漫画笔记把高频考点再过一遍，心里有底才发现其实没白学。",
        "extra": ["#考前冲刺"],
    },
]

def check(title, body):
    tc = len(title)
    bc = len(body.replace("\n", "").replace(" ", ""))
    ok = "✅" if tc <= 20 and bc <= 80 else f"❌(标题{tc},正文{bc}字)"
    return ok, tc, bc

def gen_txts():
    for post in POSTS:
        title = post["title"]
        body = post["body"]
        tags = [P_BASE] + post["extra"] + [P_TAIL]
        topics = " ".join(tags)
        account = f"14芳芳财会|{post['account']}"
        ok, tc, bc = check(title, body)
        print(f"  {ok} 标题{tc}字 正文{bc}字 | {account}")
        keyword = title[:8]
        fname = f"{post['num']}_{keyword}_{post['account']}.txt"
        path = os.path.join(BASE, fname)
        content = f"标题：{title}\n\n正文：\n{body}\n\n话题：{topics}\n\n时间：\n\n账号：{account}\n"
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        post["_fname"] = fname

def gen_word():
    doc = Document()
    for i, post in enumerate(POSTS):
        body = post["body"]
        tags = [P_BASE] + post["extra"] + [P_TAIL]
        kept_tags = []
        for group in tags:
            for tag in group.split():
                if tag not in STRIP_TAGS and tag not in kept_tags:
                    kept_tags.append(tag)
        p1 = doc.add_paragraph()
        r1 = p1.add_run(body)
        r1.font.size = Pt(12)
        p2 = doc.add_paragraph()
        r2 = p2.add_run(" ".join(kept_tags))
        r2.font.size = Pt(12)
        if i != len(POSTS) - 1:
            p3 = doc.add_paragraph()
            run = p3.add_run()
            run.add_break(WD_BREAK.PAGE)
    path = os.path.join(BASE, DOCNAME)
    doc.save(path)
    print(f"  ✅ Word汇总已保存：{path}")

def make_zip():
    zip_path = os.path.join(BASE, ZIPNAME)
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for post in POSTS:
            zf.write(os.path.join(BASE, post["_fname"]), post["_fname"])
        zf.write(os.path.join(BASE, DOCNAME), DOCNAME)
    return zip_path

if __name__ == "__main__":
    print("=== 生成9篇引流txt ===")
    gen_txts()
    print("\n=== 生成Word汇总 ===")
    gen_word()
    zp = make_zip()
    print(f"\n  ✅ ZIP打包完成：{zp}")
