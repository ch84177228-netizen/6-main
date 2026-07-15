"""
引流小号 第43-49篇生成脚本
已有主题（避免重复）：
1-7: 下班学不动/资料太多/教材想睡/没出门/没时间复习/全记错了/考前两月救吗
8-14: 不知从哪下手/错误率高崩溃/背了就忘/别人进度快/孩子太闹/碎片时间/焦虑失眠
15-21: 目录没背完87天/不敢说进度/计划表5遍没完/看课懂了做题全错/初级3年后备考/午休被发现/刷300题分数上不去
22-28: 报名费不考亏/睡前30分钟/手写笔记没用/开教材刷手机/背题坐过站/辅导班更慌/越近越不想翻书
29-35: 考前全忘了/公式背了忘/家人催问/看别人打卡慌/出差断学/看了个寂寞/模拟题错一半
36-42: 备考到一半换教材/同一科报两家机构/备忘录存100条没打开/模拟考偷看答案/考前一周发现一章没看/整理桌面一小时/差三分没过的那年
"""

import os
import zipfile
from docx import Document

BASE = "/home/user/6-main/01-6个主号内容创作/引流小号_发布文案"

articles = [
    {
        "no": 43,
        "account": "沈小辉",
        "keyword": "同一个知识点反复看了五遍还是没懂",
        "title": "同一个知识点反复看了五遍还是没懂",
        "body": (
            "有个知识点🔁\n"
            "我看第一遍：好像懂了\n"
            "第二遍：哦没懂\n"
            "第五遍：好像还是没懂😶\n"
            "\n"
            "后来换了诗雨漫画笔记\n"
            "画出来的框架一眼就清楚了"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考方法 #备考日记 #诗雨会计",
    },
    {
        "no": 44,
        "account": "孙青5263",
        "keyword": "把教材刷完一遍上模拟题全是生面孔",
        "title": "把教材刷完一遍上模拟题全是生面孔",
        "body": (
            "用了三个月刷完整本教材📖\n"
            "信心满满去做模拟题\n"
            "\n"
            "第一套：不认识😶\n"
            "第二套：还是不认识\n"
            "\n"
            "诗雨漫画笔记是按考点编的\n"
            "刷完才知道方向对了"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #刷题 #备考方法 #诗雨会计",
    },
    {
        "no": 45,
        "account": "橙子学会计",
        "keyword": "买了三套备考资料结果一套都没用完",
        "title": "买了三套备考资料结果一套都没用完",
        "body": (
            "备考第一步：买资料💸\n"
            "第一套：入门太难\n"
            "第二套：太厚了放弃\n"
            "第三套：包装都没拆过😶\n"
            "\n"
            "最后用了诗雨漫画笔记\n"
            "图多字少，拆开就能翻"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考焦虑 #备考日记 #诗雨会计",
    },
    {
        "no": 46,
        "account": "新老婆",
        "keyword": "开1.5倍速刷完全程其实啥都没记住",
        "title": "开1.5倍速刷完全程其实啥都没记住",
        "body": (
            "喜欢开1.5倍速看视频课⏩\n"
            "感觉一小时学了一个半小时\n"
            "\n"
            "直到做题才发现\n"
            "听完等于没听😶\n"
            "\n"
            "换了诗雨漫画笔记对着图自己讲\n"
            "讲得出来才算真的记住了"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考方法 #备考焦虑 #诗雨会计",
    },
    {
        "no": 47,
        "account": "孙玲小红书",
        "keyword": "同一道错题刷了三遍还是照样做错",
        "title": "同一道错题刷了三遍还是照样做错",
        "body": (
            "有道题标注了重点复习🔁\n"
            "第二次做：还是错\n"
            "第三次：依然错😮‍💨\n"
            "\n"
            "后来才明白\n"
            "这个知识点根本没理解透\n"
            "\n"
            "诗雨漫画笔记画了逻辑图\n"
            "搞懂了才没再错"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #刷题 #备考方法 #诗雨会计",
    },
    {
        "no": 48,
        "account": "爹爹",
        "keyword": "备考学到一半突然不知道自己在干什么",
        "title": "备考学到一半突然不知道自己在干什么",
        "body": (
            "有时候学着学着会发呆😶\n"
            "脑子突然空了\n"
            "不知道自己在干嘛\n"
            "\n"
            "后来拿诗雨漫画笔记翻了整体框架\n"
            "看到全貌才知道自己在哪了"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考日记 #备考焦虑 #诗雨会计",
    },
    {
        "no": 49,
        "account": "张菊香",
        "keyword": "考完走出考场觉得不错成绩出来差很多",
        "title": "考完走出考场觉得不错成绩出来差很多",
        "body": (
            "考完出来感觉还行😏\n"
            "跟同学对了几道题\n"
            "自我感觉稳了\n"
            "\n"
            "成绩出来：差了一截😶\n"
            "\n"
            "复盘发现题目是似懂非懂\n"
            "用诗雨漫画笔记重新理了框架\n"
            "明年不能再靠感觉了"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考方法 #备考焦虑 #诗雨会计",
    },
]

def make_txt(a):
    return (
        f"标题：{a['title']}\n"
        f"\n"
        f"正文：\n"
        f"{a['body']}\n"
        f"\n"
        f"话题：{a['tags']}\n"
        f"\n"
        f"时间：\n"
        f"\n"
        f"账号：14芳芳财会|{a['account']}\n"
    )

txt_paths = []
for a in articles:
    fname = f"{a['no']}_{a['keyword']}_{a['account']}.txt"
    fpath = os.path.join(BASE, fname)
    with open(fpath, "w", encoding="utf-8") as f:
        f.write(make_txt(a))
    txt_paths.append(fpath)
    print(f"✅ {fname}")

# Word document
doc = Document()
doc.add_heading("引流小号文案 第43-49篇", 0)
for a in articles:
    doc.add_heading(f"{a['no']}. 账号：14芳芳财会|{a['account']}", 1)
    p = doc.add_paragraph()
    p.add_run("标题：").bold = True
    p.add_run(a["title"])
    doc.add_paragraph()
    doc.add_paragraph("正文：")
    doc.add_paragraph(a["body"])
    doc.add_paragraph()
    p2 = doc.add_paragraph()
    p2.add_run("话题：")
    p2.add_run(a["tags"])
    doc.add_paragraph(f"账号：14芳芳财会|{a['account']}")
    doc.add_paragraph("—" * 30)

word_path = os.path.join(BASE, "引流小号_第43-49篇.docx")
doc.save(word_path)
print(f"✅ Word: {word_path}")

# ZIP（txt only）
zip_path = os.path.join(BASE, "引流小号_第43-49篇_7篇文案.zip")
with zipfile.ZipFile(zip_path, "w") as zf:
    for fp in txt_paths:
        zf.write(fp, os.path.basename(fp))
print(f"✅ ZIP: {zip_path}")
print("全部完成！")
