"""
引流小号 第29-35篇生成脚本
已有主题（避免重复）：
1-7: 下班学不动/资料太多/教材想睡/没出门/没时间复习/全记错了/考前两月救吗
8-14: 不知从哪下手/错误率高崩溃/背了就忘/别人进度快/孩子太闹/碎片时间/焦虑失眠
15-21: 目录没背完87天/不敢说进度/计划表5遍没完/看课懂了做题全错/初级3年后备考/午休被发现/刷300题分数上不去
22-28: 报名费不考亏/睡前30分钟/手写笔记没用/开教材刷手机/背题坐过站/辅导班更慌/越近越不想翻书
"""

import os
import zipfile
from docx import Document

BASE = "/home/user/6-main/01-6个主号内容创作/引流小号_发布文案"

articles = [
    {
        "no": 29,
        "account": "沈小辉",
        "keyword": "考前一个月突然全忘了",
        "title": "考前一个月突然觉得全都忘了",
        "body": (
            "距考试一个月了🗓️\n"
            "突然发现啥都想不起来\n"
            "\n"
            "明明之前看了好几遍\n"
            "但坐下来一片空白😱\n"
            "\n"
            "换成诗雨漫画笔记重新过一遍\n"
            "图解记忆，好歹找回一些感觉"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #考前冲刺 #备考焦虑 #诗雨会计",
    },
    {
        "no": 30,
        "account": "孙青5263",
        "keyword": "公式背完第二天就忘了",
        "title": "公式背完第二天就忘了咋整",
        "body": (
            "会计公式那么多，背完这个忘那个📝\n"
            "考前感觉全是新知识\n"
            "\n"
            "明明背过三遍的公式😭\n"
            "下笔就卡住了\n"
            "\n"
            "后来配诗雨漫画笔记理解着背\n"
            "推导一下，比硬背留得住"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考方法 #考前冲刺 #诗雨会计",
    },
    {
        "no": 31,
        "account": "橙子学会计",
        "keyword": "家人每天催问备考进度",
        "title": "家人每天问备考到哪了真的烦",
        "body": (
            "备考最难熬的不是刷题\n"
            "是家人每天在旁边问进度😮‍💨\n"
            "\n"
            "「看了没」「能过吗」「报个班嘛」\n"
            "我：……\n"
            "\n"
            "后来拿诗雨漫画笔记给他们看\n"
            "看起来挺努力，问少了😅"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考日记 #上班族备考 #诗雨会计",
    },
    {
        "no": 32,
        "account": "新老婆",
        "keyword": "刷到别人打卡备考开始慌",
        "title": "刷到别人备考打卡笔记开始慌了",
        "body": (
            "刷小红书刷到别人的备考打卡😶\n"
            "每天5点起，刷几百题，背几十页\n"
            "\n"
            "我：今天只翻了封面\n"
            "焦虑到直接关掉手机😩\n"
            "\n"
            "后来换成诗雨漫画笔记随便翻\n"
            "翻了就算学了，不对比了"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考焦虑 #备考日记 #诗雨会计",
    },
    {
        "no": 33,
        "account": "孙玲小红书",
        "keyword": "出差一周回来不知从哪接着学",
        "title": "出差一周回来完全不知从哪接着学",
        "body": (
            "出差一周，临行前发誓要抓紧学✈️\n"
            "结果每天到酒店已经快十二点了\n"
            "\n"
            "回来翻开教材，全陌生了😔\n"
            "这一周等于白学\n"
            "\n"
            "拿诗雨漫画笔记翻了翻\n"
            "图解版，捡起来快一点"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #上班族备考 #备考日记 #诗雨会计",
    },
    {
        "no": 34,
        "account": "爹爹",
        "keyword": "备考两月才发现全看了个寂寞",
        "title": "备考两个月才发现全看了个寂寞",
        "body": (
            "踏踏实实看了两个月教材💪\n"
            "感觉掌握得差不多了\n"
            "\n"
            "做模拟题：第一题就卡住\n"
            "全程靠蒙😶\n"
            "\n"
            "后来才知道，看书不等于会做题😵\n"
            "换诗雨漫画笔记配着习题一起走\n"
            "找到感觉了"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #刷题 #备考方法 #诗雨会计",
    },
    {
        "no": 35,
        "account": "张菊香",
        "keyword": "模拟题错一半还有必要继续吗",
        "title": "模拟题错了一半，还有必要继续吗",
        "body": (
            "上周做了套模拟题😩\n"
            "错了一半，原来的信心全碎了\n"
            "\n"
            "一度想着：要不这年别考了\n"
            "报名费都白交了也认了\n"
            "\n"
            "朋友说先别急，看看哪里错的\n"
            "拿诗雨漫画笔记对着错题捋了一遍\n"
            "感觉还有救"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #刷题 #考前冲刺 #诗雨会计",
    },
]

def make_txt(a):
    return (
        f"标题：{a['title']}\n"
        f"\n"
        f"正文：\n"
        f"{a['body']}\n"
        f"\n"
        f"话题：{a['tags']}\n"
        f"\n"
        f"时间：\n"
        f"\n"
        f"账号：14芳芳财会|{a['account']}\n"
    )

txt_paths = []
for a in articles:
    fname = f"{a['no']}_{a['keyword']}_{a['account']}.txt"
    fpath = os.path.join(BASE, fname)
    with open(fpath, "w", encoding="utf-8") as f:
        f.write(make_txt(a))
    txt_paths.append(fpath)
    print(f"✅ {fname}")

# Word document
doc = Document()
doc.add_heading("引流小号文案 第29-35篇", 0)
for a in articles:
    doc.add_heading(f"{a['no']}. 账号：14芳芳财会|{a['account']}", 1)
    p = doc.add_paragraph()
    p.add_run("标题：").bold = True
    p.add_run(a["title"])
    doc.add_paragraph()
    doc.add_paragraph("正文：")
    doc.add_paragraph(a["body"])
    doc.add_paragraph()
    doc.add_paragraph(f"话题：{a['tags']}")
    doc.add_paragraph(f"账号：14芳芳财会|{a['account']}")
    doc.add_paragraph("—" * 30)

word_path = os.path.join(BASE, "引流小号_第29-35篇.docx")
doc.save(word_path)
print(f"✅ Word: {word_path}")

# ZIP（txt only）
zip_path = os.path.join(BASE, "引流小号_第29-35篇_7篇文案.zip")
with zipfile.ZipFile(zip_path, "w") as zf:
    for fp in txt_paths:
        zf.write(fp, os.path.basename(fp))
print(f"✅ ZIP: {zip_path}")
print("全部完成！")
