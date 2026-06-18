"""
引流小号 第64-70篇生成脚本
已有主题（避免重复）：
1-7: 下班学不动/资料太多/教材想睡/没出门/没时间复习/全记错了/考前两月救吗
8-14: 不知从哪下手/错误率高崩溃/背了就忘/别人进度快/孩子太闹/碎片时间/焦虑失眠
15-21: 目录没背完87天/不敢说进度/计划表5遍没完/看课懂了做题全错/初级3年后备考/午休被发现/刷300题上不去
22-28: 报名费不考亏/睡前30分钟/手写笔记没用/开教材刷手机/背题坐过站/辅导班更慌/越近越不想翻书
29-35: 考前全忘了/公式背了忘/家人催问/看别人打卡慌/出差断学/看了个寂寞/模拟题错一半
36-42: 备考到一半换教材/同一科报两家机构/备忘录100条没打开/模拟考偷看答案/考前一周发现一章没看/整理桌面一小时/差三分没过的那年
43-49: 同一知识点看五遍没懂/教材刷完模拟题全陌生/三套资料一套没用完/1.5倍速啥没记住/同一错题刷三遍还错/学到一半不知在干啥/考完感觉稳成绩差很多
50-56: 报了网课囤着没去上/备考三个月体重涨五斤/笔记精美但不复习/纠结换证书方向/别人说简单我看了好几天/定十个闹钟没起来/备考群变聊天群
57-63: 学完新章节上一章忘了/状态最好偏偏不学/大纲列计划学不完/今年去年进度一样/知识点三遍绕不过来/就差这章明天再学/算了有时间没认真学
"""

import os
import zipfile
from docx import Document

BASE = "/home/user/6-main/01-6个主号内容创作/引流小号_发布文案"

articles = [
    {
        "no": 64,
        "account": "沈小辉",
        "keyword": "模拟题分数越练越低越练越慌",
        "title": "模拟题分数越练越低越练越慌",
        "body": (
            "练了两套模拟题📊\n"
            "第一套68分，第二套61分\n"
            "感觉越练越不会了😶\n"
            "\n"
            "后来发现是基础没打牢\n"
            "用诗雨漫画笔记把错题知识点逐个击破\n"
            "分数才慢慢稳住"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #刷题 #备考方法 #诗雨会计",
    },
    {
        "no": 65,
        "account": "孙青5263",
        "keyword": "复习完一遍发现什么都没留下",
        "title": "复习完一遍发现什么都没留下",
        "body": (
            "三月份开始备考📚\n"
            "把教材啃完了整整一遍\n"
            "坐下来做题，脑子空白😶\n"
            "\n"
            "后来用诗雨漫画笔记做复习\n"
            "每章有重点图，反复看才挂住"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考方法 #备考焦虑 #诗雨会计",
    },
    {
        "no": 66,
        "account": "橙子学会计",
        "keyword": "考前两周突然发现有整章没看",
        "title": "考前两周突然发现有整章没看",
        "body": (
            "考前两周翻了下目录📋\n"
            "发现有一整章完全没碰\n"
            "连名字都觉得陌生😶\n"
            "\n"
            "临时抱佛脚，把诗雨漫画笔记翻出来\n"
            "两天啃完这章，勉强能作答"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考日记 #备考焦虑 #诗雨会计",
    },
    {
        "no": 67,
        "account": "新老婆",
        "keyword": "越备考越觉得自己什么都不会",
        "title": "越备考越觉得自己什么都不会",
        "body": (
            "备考前觉得自己基础还行🙂\n"
            "备考两个月后发现\n"
            "自己其实什么都不会😶\n"
            "\n"
            "后来按诗雨漫画笔记重新梳理框架\n"
            "发现漏的全是考点，不是不行是没系统"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考焦虑 #备考方法 #诗雨会计",
    },
    {
        "no": 68,
        "account": "孙玲小红书",
        "keyword": "备考中途被领导发现偷偷请假",
        "title": "备考中途被领导发现偷偷请假",
        "body": (
            "领导问为什么最近总请假😅\n"
            "犹豫了两秒说「陪家人」\n"
            "其实在家备考中级会计😶\n"
            "\n"
            "后来买了诗雨漫画笔记随身带\n"
            "路上、饭点都能翻，不用再请假了"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考日记 #打工人 #诗雨会计",
    },
    {
        "no": 69,
        "account": "爹爹",
        "keyword": "手机一弹通知就放下教材了",
        "title": "手机一弹通知就放下教材了",
        "body": (
            "备考最大的敌人不是考题📵\n"
            "是桌上那个亮屏的手机\n"
            "\n"
            "每次「叮」一声，书就合上了😶\n"
            "\n"
            "后来把诗雨漫画笔记放手机旁边\n"
            "亮屏的时候顺手翻一页"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考日记 #备考方法 #诗雨会计",
    },
    {
        "no": 70,
        "account": "张菊香",
        "keyword": "每次刷完题都觉得记住了下次又不会",
        "title": "每次刷完题觉得记住了下次又不会",
        "body": (
            "刷题的时候对完答案✔️\n"
            "感觉这道题学会了\n"
            "下次遇到同类题，还是不会😶\n"
            "\n"
            "后来用诗雨漫画笔记把同类题归了归\n"
            "系统理解才是真的会"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #刷题 #备考方法 #诗雨会计",
    },
]

def make_txt(a):
    return (
        f"标题：{a['title']}\n"
        f"\n"
        f"正文：\n"
        f"{a['body']}\n"
        f"\n"
        f"话题：{a['tags']}\n"
        f"\n"
        f"时间：\n"
        f"\n"
        f"账号：14芳芳财会|{a['account']}\n"
    )

txt_paths = []
for a in articles:
    fname = f"{a['no']}_{a['keyword']}_{a['account']}.txt"
    fpath = os.path.join(BASE, fname)
    with open(fpath, "w", encoding="utf-8") as f:
        f.write(make_txt(a))
    txt_paths.append(fpath)
    body_len = len(a["body"].replace("\n", ""))
    print(f"✅ {fname} (正文{body_len}字)")

# Word document
doc = Document()
doc.add_heading("引流小号文案 第64-70篇", 0)
for a in articles:
    doc.add_heading(f"{a['no']}. 账号：14芳芳财会|{a['account']}", 1)
    p = doc.add_paragraph()
    p.add_run("标题：").bold = True
    p.add_run(a["title"])
    doc.add_paragraph()
    doc.add_paragraph("正文：")
    doc.add_paragraph(a["body"])
    doc.add_paragraph()
    p2 = doc.add_paragraph()
    p2.add_run("话题：")
    p2.add_run(a["tags"])
    doc.add_paragraph(f"账号：14芳芳财会|{a['account']}")
    doc.add_paragraph("—" * 30)

word_path = os.path.join(BASE, "引流小号_第64-70篇.docx")
doc.save(word_path)
print(f"✅ Word: {word_path}")

# ZIP（txt only）
zip_path = os.path.join(BASE, "引流小号_第64-70篇_7篇文案.zip")
with zipfile.ZipFile(zip_path, "w") as zf:
    for fp in txt_paths:
        zf.write(fp, os.path.basename(fp))
print(f"✅ ZIP: {zip_path}")
print("全部完成！")
