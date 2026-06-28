#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import os, zipfile
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_BREAK

BASE = "/home/user/6-main/01-6个主号内容创作/引流小号_发布文案"
P = "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考日记 #上班族备考 #考前冲刺 #诗雨会计"
DROP_TAGS = {"#财会", "#备考日记", "#上班族备考", "#考前冲刺"}

def save(path, content):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)

records = []

def write_txt(idx, title, body, account_name):
    tc = len(title)
    bc = len(body.replace("\n","").replace(" ",""))
    ok = "✅" if tc<=20 and bc<=80 else f"❌(标题{tc},正文{bc}字)"
    print(f"  {ok} 标题{tc}字 正文{bc}字 | {account_name}")
    kw = title[:8].replace("？","").replace("！","").replace("，","").replace("。","")
    fname = f"{idx}_{kw}_{account_name}.txt"
    content = f"标题：{title}\n\n正文：\n{body}\n\n话题：{P}\n\n时间：\n\n账号：14芳芳财会|{account_name}\n"
    save(os.path.join(BASE, fname), content)
    records.append((idx, body))

def gen_word(start, end):
    doc = Document()
    first = True
    for idx, body in records:
        if not first:
            p = doc.add_paragraph()
            p.add_run().add_break(WD_BREAK.PAGE)
        h = doc.add_paragraph()
        hr = h.add_run(f"第{idx}篇")
        hr.bold = True
        hr.font.size = Pt(14)
        for line in body.strip().split("\n"):
            p = doc.add_paragraph()
            r = p.add_run(line)
            r.font.size = Pt(12)
        tags = [t for t in P.split() if t not in DROP_TAGS]
        p = doc.add_paragraph()
        r = p.add_run(" ".join(tags))
        r.font.size = Pt(12)
        first = False
    out = os.path.join(BASE, f"引流小号_第{start}-{end}篇.docx")
    doc.save(out)
    print(f"  ✅ 汇总Word：{out}")

def make_zip(start, end):
    zip_path = os.path.join(BASE, f"引流小号_第{start}-{end}篇.zip")
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for i in range(start, end+1):
            for fname in os.listdir(BASE):
                if fname.startswith(f"{i}_") and fname.endswith(".txt"):
                    zf.write(os.path.join(BASE, fname), fname)
        wname = f"引流小号_第{start}-{end}篇.docx"
        if os.path.exists(os.path.join(BASE, wname)):
            zf.write(os.path.join(BASE, wname), wname)
    return zip_path

articles = [
    (133, "综合题一上来就不知道从哪下手",
     "做综合题最大的问题是没思路，题目一长就懵，不知道先看哪一步。😵后来用诗雨漫画笔记把每类题的解题顺序记下来，照着走，至少不空白了。",
     "沈小辉"),
    (134, "时段法确认收入老是算错进度",
     "建造合同按进度确认收入，履约进度怎么算我总搞混。🤯翻了诗雨漫画笔记，已投入成本除以预计总成本，一下就清楚了，算进度不再乱。",
     "孙青5263"),
    (135, "下班只想躺平根本不想碰书",
     "上一天班回家只想瘫着，书摆在那一周没翻过。😮‍💨后来给自己定每天只看五页漫画笔记，门槛低了反而坚持下来了，慢慢有了状态。",
     "橙子学会计"),
    (136, "金融资产减值那块完全没头绪",
     "预期信用损失三阶段，12个月和整个存续期老是分不清。😣看了诗雨漫画笔记的图解，三阶段一张图摆清楚，终于不糊涂了。",
     "新老婆"),
    (137, "刷了好多题分数还是没起色",
     "题刷了一本又一本，模考分数还是上不去，开始怀疑方法不对。😔后来发现是没抓重点。诗雨漫画笔记把高频考点标得很清楚，刷题更有方向。",
     "孙玲小红书"),
    (138, "看着倒计时心跳都快了",
     "每次看到倒计时还剩两个多月，心就一紧，越急越学不进去。😰后来用诗雨漫画笔记一天稳稳搞懂一个考点，焦虑反而减轻了。",
     "爹爹"),
    (139, "教材太厚翻开就犯困",
     "中级实务那本教材又厚又枯燥，翻开没几页就开始打瞌睡。😴换成诗雨漫画笔记，图多有意思，居然能看进去，不那么催眠了。",
     "张菊香"),
    (140, "学了后面忘前面像狗熊掰棒子",
     "学到后面章节，前面学的全忘了，像狗熊掰棒子一样。😭后来用诗雨漫画笔记定期回看图，记忆更牢，前后串起来了。",
     "孙文礼"),
    (141, "怕考不过又不甘心放弃",
     "一边怕自己考不过白忙一场，一边又不甘心放弃，特别纠结。😣后来想通了，与其纠结不如行动，每天靠诗雨漫画笔记搞懂一点，踏实多了。",
     "晴天"),
]

if __name__ == "__main__":
    print("=== 生成引流小号第133-141篇 ===")
    for idx, title, body, account in articles:
        write_txt(idx, title, body, account)
    gen_word(133, 141)
    zp = make_zip(133, 141)
    print(f"\n  ✅ ZIP打包完成：{zp}")
