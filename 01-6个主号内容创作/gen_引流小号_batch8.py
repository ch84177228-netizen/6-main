"""
引流小号 第57-63篇生成脚本
已有主题（避免重复）：
1-7: 下班学不动/资料太多/教材想睡/没出门/没时间复习/全记错了/考前两月救吗
8-14: 不知从哪下手/错误率高崩溃/背了就忘/别人进度快/孩子太闹/碎片时间/焦虑失眠
15-21: 目录没背完87天/不敢说进度/计划表5遍没完/看课懂了做题全错/初级3年后备考/午休被发现/刷300题分数上不去
22-28: 报名费不考亏/睡前30分钟/手写笔记没用/开教材刷手机/背题坐过站/辅导班更慌/越近越不想翻书
29-35: 考前全忘了/公式背了忘/家人催问/看别人打卡慌/出差断学/看了个寂寞/模拟题错一半
36-42: 备考到一半换教材/同一科报两家机构/备忘录存100条没打开/模拟考偷看答案/考前一周发现一章没看/整理桌面一小时/差三分没过的那年
43-49: 同一知识点看五遍没懂/教材刷完模拟题全陌生/三套资料一套没用完/1.5倍速啥没记住/同一错题刷三遍还错/学到一半不知在干啥/考完感觉稳成绩差很多
50-56: 报了网课囤着没去上/备考三个月体重涨五斤/笔记精美但不复习/纠结换证书方向/别人说简单我看了好几天/定十个闹钟没起来/备考群变聊天群
"""

import os
import zipfile
from docx import Document

BASE = "/home/user/6-main/01-6个主号内容创作/引流小号_发布文案"

articles = [
    {
        "no": 57,
        "account": "沈小辉",
        "keyword": "学完新章节才发现上一章全忘了",
        "title": "学完新章节才发现上一章全忘了",
        "body": (
            "好不容易搞懂了第五章🎉\n"
            "翻到第六章继续学\n"
            "\n"
            "做完一道题回头要用到第四章\n"
            "脑子一片空白😶\n"
            "\n"
            "后来用诗雨漫画笔记把框架拎起来\n"
            "前后关联才挂得住"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考方法 #备考日记 #诗雨会计",
    },
    {
        "no": 58,
        "account": "孙青5263",
        "keyword": "备考状态最好那天我偏偏不在学习",
        "title": "备考状态最好那天我偏偏不在学习",
        "body": (
            "备考最认真的时候是在饭桌上🍜\n"
            "脑子里复盘考点，思路特别清晰\n"
            "\n"
            "真坐下来打开教材\n"
            "突然什么都不想了😶\n"
            "\n"
            "后来把诗雨漫画笔记放饭桌旁边\n"
            "吃饭时翻两页，也算学了"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考方法 #备考日记 #诗雨会计",
    },
    {
        "no": 59,
        "account": "橙子学会计",
        "keyword": "用考试大纲列学习计划发现根本学不完",
        "title": "用考试大纲列学习计划发现根本学不完",
        "body": (
            "备考第一步：下载考试大纲📋\n"
            "按大纲列了一份完整计划表\n"
            "\n"
            "结果发现每天要学好几章\n"
            "三个月根本不够用😶\n"
            "\n"
            "后来用诗雨漫画笔记找到高频考点\n"
            "先学最重要的，学一章稳一章"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考方法 #备考焦虑 #诗雨会计",
    },
    {
        "no": 60,
        "account": "新老婆",
        "keyword": "今年和去年备考进度简直一模一样",
        "title": "今年和去年备考进度简直一模一样",
        "body": (
            "翻了翻去年的备考记录🗓️\n"
            "6月：刚看完第三章\n"
            "今年6月：刚看完第三章\n"
            "\n"
            "……😶\n"
            "\n"
            "今年换了诗雨漫画笔记重来\n"
            "图解版看起来快一倍，不能再一样了"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考日记 #备考焦虑 #诗雨会计",
    },
    {
        "no": 61,
        "account": "孙玲小红书",
        "keyword": "这个知识点学了三遍还是绕不过来",
        "title": "这个知识点学了三遍还是绕不过来",
        "body": (
            "投资收益的确认时点🙃\n"
            "学了一遍：好像懂了\n"
            "复习第二遍：又蒙了\n"
            "刷到题目：完全不会用😶\n"
            "\n"
            "后来用诗雨漫画笔记找到时间轴图解\n"
            "对着图捋一遍才真的会了"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考方法 #刷题 #诗雨会计",
    },
    {
        "no": 62,
        "account": "爹爹",
        "keyword": "一直说就差这一章等明天再开始学",
        "title": "一直说就差这一章等明天再开始学",
        "body": (
            "备考每周的口头禅📅\n"
            "「就差这一章，今天先歇」\n"
            "「还有这一章，明天学」\n"
            "「马上了马上了」😶\n"
            "\n"
            "后来把诗雨漫画笔记摊在桌上\n"
            "顺手就翻了，不用再找理由"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考日记 #备考焦虑 #诗雨会计",
    },
    {
        "no": 63,
        "account": "张菊香",
        "keyword": "算了下距考试还有时间就没认真学",
        "title": "算了下距考试还有时间就没认真学",
        "body": (
            "5月报完名算了一下📆\n"
            "距考试还有4个月\n"
            "感觉很够用\n"
            "\n"
            "7月才开始慌：只剩50天了😶\n"
            "\n"
            "后来碰到这种感觉就拿诗雨漫画笔记翻一页\n"
            "至少走一步，别算了"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考焦虑 #备考方法 #诗雨会计",
    },
]

def make_txt(a):
    return (
        f"标题：{a['title']}\n"
        f"\n"
        f"正文：\n"
        f"{a['body']}\n"
        f"\n"
        f"话题：{a['tags']}\n"
        f"\n"
        f"时间：\n"
        f"\n"
        f"账号：14芳芳财会|{a['account']}\n"
    )

txt_paths = []
for a in articles:
    fname = f"{a['no']}_{a['keyword']}_{a['account']}.txt"
    fpath = os.path.join(BASE, fname)
    with open(fpath, "w", encoding="utf-8") as f:
        f.write(make_txt(a))
    txt_paths.append(fpath)
    body_len = len(a["body"].replace("\n", ""))
    print(f"✅ {fname} (正文{body_len}字)")

# Word document
doc = Document()
doc.add_heading("引流小号文案 第57-63篇", 0)
for a in articles:
    doc.add_heading(f"{a['no']}. 账号：14芳芳财会|{a['account']}", 1)
    p = doc.add_paragraph()
    p.add_run("标题：").bold = True
    p.add_run(a["title"])
    doc.add_paragraph()
    doc.add_paragraph("正文：")
    doc.add_paragraph(a["body"])
    doc.add_paragraph()
    p2 = doc.add_paragraph()
    p2.add_run("话题：")
    p2.add_run(a["tags"])
    doc.add_paragraph(f"账号：14芳芳财会|{a['account']}")
    doc.add_paragraph("—" * 30)

word_path = os.path.join(BASE, "引流小号_第57-63篇.docx")
doc.save(word_path)
print(f"✅ Word: {word_path}")

# ZIP（txt only）
zip_path = os.path.join(BASE, "引流小号_第57-63篇_7篇文案.zip")
with zipfile.ZipFile(zip_path, "w") as zf:
    for fp in txt_paths:
        zf.write(fp, os.path.basename(fp))
print(f"✅ ZIP: {zip_path}")
print("全部完成！")
