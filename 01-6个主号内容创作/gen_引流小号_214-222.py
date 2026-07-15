#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import os, zipfile
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_BREAK

BASE = "/home/user/6-main/01-6个主号内容创作/引流小号_发布文案"
ZIPNAME = "编号20260709-07_引流账号文案_芳芳财会.zip"
DOCNAME = "引流小号_第214-222篇.docx"
P_BASE = "#会计 #中级会计 #中级会计备考 #注会cpa #财会"
P_TAIL = "#诗雨会计"
STRIP_TAGS = ["#财会", "#备考日记", "#上班族备考", "#考前冲刺", "#备考焦虑"]

POSTS = [
    {
        "num": 214, "account": "沈小辉",
        "title": "报名后才发现自己时间根本不够",
        "body": "一时冲动报了名，真开始学才发现每天挤不出多少时间。😮‍💨后来靠诗雨漫画笔记利用零碎时间翻两页，积少成多也追得上。",
        "extra": ["#上班族备考"],
    },
    {
        "num": 215, "account": "孙青5263",
        "title": "越想学好越不敢开始",
        "body": "总想着要么不学、要学就一次学扎实，结果反而迟迟不敢开始。😅后来告诉自己先翻两页诗雨漫画笔记，开了头就没那么难了。",
        "extra": ["#会计学习方法"],
    },
    {
        "num": 216, "account": "橙子学会计",
        "title": "听课全懂做题全废",
        "body": "跟着老师听课节节都懂，一自己动手做题就全废。😩后来靠诗雨漫画笔记把考点自己复述一遍，才知道哪里是真懂。",
        "extra": ["#中级会计实务"],
    },
    {
        "num": 217, "account": "新老婆",
        "title": "笔记记得密密麻麻却用不上",
        "body": "笔记记得满满一本，密密麻麻，真复习时反而找不到重点。😵后来换成诗雨漫画笔记，高频考点一目了然，复习不再抓瞎。",
        "extra": ["#会计学习方法"],
    },
    {
        "num": 218, "account": "孙玲小红书",
        "title": "背了忘忘了背来回折腾",
        "body": "同一个知识点背了忘、忘了背，来回好几遍还是记不牢。😤后来靠诗雨漫画笔记靠图理解着记，比死记硬背牢多了。",
        "extra": ["#中级会计备考"],
    },
    {
        "num": 219, "account": "爹爹",
        "title": "下班累到只想瘫着不想动脑",
        "body": "上一天班回家整个人累瘫，让我动脑看书实在提不起劲。😴后来靠诗雨漫画笔记，图多字少不费脑，瘫着也能看两眼。",
        "extra": ["#上班族备考"],
    },
    {
        "num": 220, "account": "张菊香",
        "title": "一看到大段文字就头大",
        "body": "教材一大段一大段的文字，看着就头大、根本读不进去。😮‍💨后来发现诗雨漫画笔记把知识点拆成图，看着轻松还记得住。",
        "extra": ["#会计学习方法"],
    },
    {
        "num": 221, "account": "孙文礼",
        "title": "计划总排得满执行却打折",
        "body": "计划表排得满满当当，真到执行永远打对折。😅后来不硬凑计划了，跟着诗雨漫画笔记有空就翻，反而更稳。",
        "extra": ["#中级会计备考"],
    },
    {
        "num": 222, "account": "晴天",
        "title": "越临近考试越静不下心",
        "body": "越接近考试心里越乱，坐下来也静不下心复习。😮‍💨后来靠诗雨漫画笔记把高频考点一遍遍过，心里有底才踏实。",
        "extra": ["#考前冲刺"],
    },
]

def check(title, body):
    tc = len(title)
    bc = len(body.replace("\n", "").replace(" ", ""))
    ok = "✅" if tc <= 20 and bc <= 80 else f"❌(标题{tc},正文{bc}字)"
    return ok, tc, bc

def gen_txts():
    for post in POSTS:
        title = post["title"]
        body = post["body"]
        tags = [P_BASE] + post["extra"] + [P_TAIL]
        topics = " ".join(tags)
        account = f"14芳芳财会|{post['account']}"
        ok, tc, bc = check(title, body)
        print(f"  {ok} 标题{tc}字 正文{bc}字 | {account}")
        keyword = title[:8]
        fname = f"{post['num']}_{keyword}_{post['account']}.txt"
        path = os.path.join(BASE, fname)
        content = f"标题：{title}\n\n正文：\n{body}\n\n话题：{topics}\n\n时间：\n\n账号：{account}\n"
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        post["_fname"] = fname

def gen_word():
    doc = Document()
    for i, post in enumerate(POSTS):
        body = post["body"]
        tags = [P_BASE] + post["extra"] + [P_TAIL]
        kept_tags = []
        for group in tags:
            for tag in group.split():
                if tag not in STRIP_TAGS and tag not in kept_tags:
                    kept_tags.append(tag)
        p1 = doc.add_paragraph()
        r1 = p1.add_run(body)
        r1.font.size = Pt(12)
        p2 = doc.add_paragraph()
        r2 = p2.add_run(" ".join(kept_tags))
        r2.font.size = Pt(12)
        if i != len(POSTS) - 1:
            p3 = doc.add_paragraph()
            run = p3.add_run()
            run.add_break(WD_BREAK.PAGE)
    path = os.path.join(BASE, DOCNAME)
    doc.save(path)
    print(f"  ✅ Word汇总已保存：{path}")

def make_zip():
    zip_path = os.path.join(BASE, ZIPNAME)
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for post in POSTS:
            zf.write(os.path.join(BASE, post["_fname"]), post["_fname"])
        zf.write(os.path.join(BASE, DOCNAME), DOCNAME)
    return zip_path

if __name__ == "__main__":
    print("=== 生成9篇引流txt ===")
    gen_txts()
    print("\n=== 生成Word汇总 ===")
    gen_word()
    zp = make_zip()
    print(f"\n  ✅ ZIP打包完成：{zp}")
