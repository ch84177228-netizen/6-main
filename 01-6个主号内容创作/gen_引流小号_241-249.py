#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import os, zipfile
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_BREAK

BASE = "/home/user/6-main/01-6个主号内容创作/引流小号_发布文案"
ZIPNAME = "编号20260715-01_引流账号文案_芳芳财会.zip"
DOCNAME = "引流小号_第241-249篇.docx"
P_BASE = "#会计 #中级会计 #中级会计备考 #注会cpa #财会"
P_TAIL = "#诗雨会计"
# 本批要求：Word 里的话题固定保留 5 个
WORD_TAGS = ["#会计", "#中级会计", "#中级会计备考", "#注会cpa", "#诗雨会计"]

POSTS = [
    {
        "num": 241, "account": "沈小辉",
        "title": "教材翻开第一页就想合上",
        "body": "教材一翻开满页密密麻麻的字，还没看两行就想合上。😮‍💨后来换成诗雨漫画笔记，图多字少，翻开反而愿意多看两页。",
        "extra": ["#会计学习方法"],
    },
    {
        "num": 242, "account": "孙青5263",
        "title": "上班太忙学习总断断续续",
        "body": "工作一忙起来学习就断了，好不容易捡起来又忘得差不多。😮‍💨后来靠诗雨漫画笔记随时翻两页，断了也能快速接上。",
        "extra": ["#上班族备考"],
    },
    {
        "num": 243, "account": "橙子学会计",
        "title": "一算题就手忙脚乱没头绪",
        "body": "一遇到计算题就手忙脚乱，公式往哪套都不知道。😵后来靠诗雨漫画笔记把公式和例子对着看，思路清楚多了。",
        "extra": ["#中级会计实务"],
    },
    {
        "num": 244, "account": "新老婆",
        "title": "零基础怕自己根本学不会",
        "body": "完全零基础，心里总打鼓怕自己压根学不会。😔后来跟着诗雨漫画笔记从最基础的看起，一点点也能入门。",
        "extra": ["#零基础备考"],
    },
    {
        "num": 245, "account": "孙玲小红书",
        "title": "报了名却一直没进入状态",
        "body": "名报了、书也买了，可就是迟迟进入不了学习状态。😅后来先翻两页诗雨漫画笔记热热身，慢慢就找回感觉了。",
        "extra": ["#会计学习方法"],
    },
    {
        "num": 246, "account": "爹爹",
        "title": "下班只想休息一学就困",
        "body": "上班一整天，下班只想歇着，一坐下学习就犯困。😴后来靠诗雨漫画笔记，图文轻松不烧脑，累的时候也翻得动。",
        "extra": ["#上班族备考"],
    },
    {
        "num": 247, "account": "张菊香",
        "title": "背了后面忘前面来回打转",
        "body": "背到后面章节，前面学过的又忘光了，来回打转很挫败。😤后来常翻诗雨漫画笔记回顾框架，前后串起来记得牢。",
        "extra": ["#中级会计备考"],
    },
    {
        "num": 248, "account": "孙文礼",
        "title": "看着别人进度快自己更慌",
        "body": "刷到别人复习进度飞快，一对比自己更慌更乱。😣后来不比了，跟着诗雨漫画笔记按自己节奏走，反而更稳。",
        "extra": ["#会计学习方法"],
    },
    {
        "num": 249, "account": "晴天",
        "title": "越到考前越怕漏了考点",
        "body": "越接近考试越心慌，总怕有高频考点没复习到。😮‍💨后来对照诗雨漫画笔记的高频考点查漏补缺，心里才有底。",
        "extra": ["#考前冲刺"],
    },
]

def check(title, body):
    tc = len(title)
    bc = len(body.replace("\n", "").replace(" ", ""))
    ok = "✅" if tc <= 20 and bc <= 80 else f"❌(标题{tc},正文{bc}字)"
    return ok, tc, bc

def gen_txts():
    for post in POSTS:
        title = post["title"]; body = post["body"]
        tags = [P_BASE] + post["extra"] + [P_TAIL]
        topics = " ".join(tags)
        account = f"14芳芳财会|{post['account']}"
        ok, tc, bc = check(title, body)
        print(f"  {ok} 标题{tc}字 正文{bc}字 | {account}")
        fname = f"{post['num']}_{title[:8]}_{post['account']}.txt"
        with open(os.path.join(BASE, fname), "w", encoding="utf-8") as f:
            f.write(f"标题：{title}\n\n正文：\n{body}\n\n话题：{topics}\n\n时间：\n\n账号：{account}\n")
        post["_fname"] = fname

def gen_word():
    doc = Document()
    tag_line = " ".join(WORD_TAGS)  # 固定5个话题
    for i, post in enumerate(POSTS):
        p1 = doc.add_paragraph(); r1 = p1.add_run(post["body"]); r1.font.size = Pt(12)
        p2 = doc.add_paragraph(); r2 = p2.add_run(tag_line); r2.font.size = Pt(12)
        if i != len(POSTS) - 1:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.save(os.path.join(BASE, DOCNAME))
    print(f"  ✅ Word汇总已保存（每篇话题固定5个）：{os.path.join(BASE, DOCNAME)}")

def make_zip():
    zp = os.path.join(BASE, ZIPNAME)
    with zipfile.ZipFile(zp, "w", zipfile.ZIP_DEFLATED) as zf:
        for post in POSTS:
            zf.write(os.path.join(BASE, post["_fname"]), post["_fname"])
        zf.write(os.path.join(BASE, DOCNAME), DOCNAME)
    return zp

if __name__ == "__main__":
    print("=== 生成9篇引流txt ===")
    gen_txts()
    print("\n=== 生成Word汇总 ===")
    gen_word()
    print(f"\n  ✅ ZIP打包完成：{make_zip()}")
