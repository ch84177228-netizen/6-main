"""
引流小号 第50-56篇生成脚本
已有主题（避免重复）：
1-7: 下班学不动/资料太多/教材想睡/没出门/没时间复习/全记错了/考前两月救吗
8-14: 不知从哪下手/错误率高崩溃/背了就忘/别人进度快/孩子太闹/碎片时间/焦虑失眠
15-21: 目录没背完87天/不敢说进度/计划表5遍没完/看课懂了做题全错/初级3年后备考/午休被发现/刷300题分数上不去
22-28: 报名费不考亏/睡前30分钟/手写笔记没用/开教材刷手机/背题坐过站/辅导班更慌/越近越不想翻书
29-35: 考前全忘了/公式背了忘/家人催问/看别人打卡慌/出差断学/看了个寂寞/模拟题错一半
36-42: 备考到一半换教材/同一科报两家机构/备忘录存100条没打开/模拟考偷看答案/考前一周发现一章没看/整理桌面一小时/差三分没过的那年
43-49: 同一知识点看五遍没懂/教材刷完模拟题全陌生/三套资料一套没用完/1.5倍速啥没记住/同一错题刷三遍还错/学到一半不知在干啥/考完感觉稳成绩差很多
"""

import os
import zipfile
from docx import Document

BASE = "/home/user/6-main/01-6个主号内容创作/引流小号_发布文案"

articles = [
    {
        "no": 50,
        "account": "沈小辉",
        "keyword": "报了网课但一直囤着没去上",
        "title": "报了网课但一直囤着没去上",
        "body": (
            "信心满满报了网课📦\n"
            "课程一直在列表里放着\n"
            "进度0%，一节没看😶\n"
            "\n"
            "后来翻开诗雨漫画笔记\n"
            "不用打开APP，直接看就能学"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考方法 #备考日记 #诗雨会计",
    },
    {
        "no": 51,
        "account": "孙青5263",
        "keyword": "备考三个月体重涨了五斤",
        "title": "备考三个月体重涨了五斤",
        "body": (
            "备考三个月⚖️\n"
            "体重涨了五斤\n"
            "朋友说：你不是在学习吗\n"
            "我：嗯，学习的同时也在吃😅\n"
            "\n"
            "后来换成诗雨漫画笔记翻着学\n"
            "没那么焦虑了，零食也吃少了"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考日记 #备考焦虑 #诗雨会计",
    },
    {
        "no": 52,
        "account": "橙子学会计",
        "keyword": "笔记做得很精美但只做了不复习",
        "title": "笔记做得很精美但只做了不复习",
        "body": (
            "笔记本买了好几本📓\n"
            "每一页都画得很精致\n"
            "配色、贴纸、手绘小图\n"
            "\n"
            "做完就没翻过第二遍😶\n"
            "\n"
            "后来用诗雨漫画笔记\n"
            "不用自己画，省下时间真复习"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考方法 #备考日记 #诗雨会计",
    },
    {
        "no": 53,
        "account": "新老婆",
        "keyword": "学到一半纠结要不要换个证书考",
        "title": "学到一半纠结要不要换个证书考",
        "body": (
            "学着学着开始纠结🤔\n"
            "要不要换个证书考\n"
            "听说那个更好就业\n"
            "\n"
            "纠结了一周，啥都没学进去😶\n"
            "\n"
            "回头还是把这本学完吧\n"
            "诗雨漫画笔记翻得快，学完心里有底"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考焦虑 #备考方法 #诗雨会计",
    },
    {
        "no": 54,
        "account": "孙玲小红书",
        "keyword": "别人说简单的章节我看了好几天",
        "title": "别人说简单的章节我看了好几天",
        "body": (
            "群里都说这章很简单✅\n"
            "我看了一天没懂\n"
            "第二天接着看，还没懂\n"
            "第三天：开始怀疑自己😮‍💨\n"
            "\n"
            "后来换成诗雨漫画笔记看图\n"
            "原来这章真没那么难"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考方法 #备考焦虑 #诗雨会计",
    },
    {
        "no": 55,
        "account": "爹爹",
        "keyword": "定了十个闹钟提醒学习一个没起来",
        "title": "定了十个闹钟提醒学习一个没起来",
        "body": (
            "怕起不来特意定了十个闹钟⏰\n"
            "第一个：关掉接着睡\n"
            "第十个：直接睡到自然醒😴\n"
            "\n"
            "后来把诗雨漫画笔记放枕头边\n"
            "醒了眼睛一睁就能翻两页"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考日记 #备考方法 #诗雨会计",
    },
    {
        "no": 56,
        "account": "张菊香",
        "keyword": "拉了备考群最后变成纯聊天群",
        "title": "拉了备考群最后变成纯聊天群",
        "body": (
            "组建了备考群互相监督📣\n"
            "第一天：打卡积极\n"
            "第一周：开始聊八卦\n"
            "现在：纯水群，没人提学习😅\n"
            "\n"
            "自己倒是没断\n"
            "配着诗雨漫画笔记自己学完了"
        ),
        "tags": "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考日记 #备考焦虑 #诗雨会计",
    },
]

def make_txt(a):
    return (
        f"标题：{a['title']}\n"
        f"\n"
        f"正文：\n"
        f"{a['body']}\n"
        f"\n"
        f"话题：{a['tags']}\n"
        f"\n"
        f"时间：\n"
        f"\n"
        f"账号：14芳芳财会|{a['account']}\n"
    )

txt_paths = []
for a in articles:
    fname = f"{a['no']}_{a['keyword']}_{a['account']}.txt"
    fpath = os.path.join(BASE, fname)
    with open(fpath, "w", encoding="utf-8") as f:
        f.write(make_txt(a))
    txt_paths.append(fpath)
    body_len = len(a["body"].replace("\n", ""))
    print(f"✅ {fname} (正文{body_len}字)")

# Word document
doc = Document()
doc.add_heading("引流小号文案 第50-56篇", 0)
for a in articles:
    doc.add_heading(f"{a['no']}. 账号：14芳芳财会|{a['account']}", 1)
    p = doc.add_paragraph()
    p.add_run("标题：").bold = True
    p.add_run(a["title"])
    doc.add_paragraph()
    doc.add_paragraph("正文：")
    doc.add_paragraph(a["body"])
    doc.add_paragraph()
    p2 = doc.add_paragraph()
    p2.add_run("话题：")
    p2.add_run(a["tags"])
    doc.add_paragraph(f"账号：14芳芳财会|{a['account']}")
    doc.add_paragraph("—" * 30)

word_path = os.path.join(BASE, "引流小号_第50-56篇.docx")
doc.save(word_path)
print(f"✅ Word: {word_path}")

# ZIP（txt only）
zip_path = os.path.join(BASE, "引流小号_第50-56篇_7篇文案.zip")
with zipfile.ZipFile(zip_path, "w") as zf:
    for fp in txt_paths:
        zf.write(fp, os.path.basename(fp))
print(f"✅ ZIP: {zip_path}")
print("全部完成！")
