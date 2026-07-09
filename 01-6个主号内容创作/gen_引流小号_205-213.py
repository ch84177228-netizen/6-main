#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import os, zipfile
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_BREAK

BASE = "/home/user/6-main/01-6个主号内容创作/引流小号_发布文案"
ZIPNAME = "编号20260709-01_引流账号文案_芳芳财会.zip"
DOCNAME = "引流小号_第205-213篇.docx"
P_BASE = "#会计 #中级会计 #中级会计备考 #注会cpa #财会"
P_TAIL = "#诗雨会计"
STRIP_TAGS = ["#财会", "#备考日记", "#上班族备考", "#考前冲刺", "#备考焦虑"]

POSTS = [
    {
        "num": 205, "account": "沈小辉",
        "title": "报名后才发现基础比想象差",
        "body": "报完名一学才发现，自己底子比想的还薄，好多名词都看不懂。😅后来从诗雨漫画笔记的基础部分补起，一点点跟上，没那么慌了。",
        "extra": ["#零基础备考"],
    },
    {
        "num": 206, "account": "孙青5263",
        "title": "每天计划学三小时实际半小时",
        "body": "计划表写着每天学三小时，真正坐下来学的可能就半小时。😮‍💨后来不定死时间了，靠诗雨漫画笔记有空翻两页，反而积累得更多。",
        "extra": ["#会计学习方法"],
    },
    {
        "num": 207, "account": "橙子学会计",
        "title": "越难的章节越想往后拖",
        "body": "遇到难懂的章节就本能地想跳过、往后拖，结果越堆越多。😩后来靠诗雨漫画笔记把难点拆开看，硬骨头也能一点点啃下来。",
        "extra": ["#中级会计实务"],
    },
    {
        "num": 208, "account": "新老婆",
        "title": "看书两小时其实在走神",
        "body": "坐在书桌前两小时，回过神来发现一大半时间都在走神发呆。😴后来用诗雨漫画笔记，图文一页一页翻，注意力反而更集中。",
        "extra": ["#会计学习方法"],
    },
    {
        "num": 209, "account": "孙玲小红书",
        "title": "背了公式却不知道怎么用",
        "body": "公式背得滚瓜烂熟，一到题目里就不知道往哪套。😣后来靠诗雨漫画笔记看公式怎么结合例子用，会背也会用了。",
        "extra": ["#中级会计实务"],
    },
    {
        "num": 210, "account": "爹爹",
        "title": "白天忙到晚只能挤晚上学",
        "body": "白天工作家庭连轴转，能学习的只剩下睡前那点时间。😮‍💨睡前翻两页诗雨漫画笔记不费脑，也算每天有进度。",
        "extra": ["#上班族备考"],
    },
    {
        "num": 211, "account": "张菊香",
        "title": "错题订正完下次照样错",
        "body": "错题认真订正了、也标了记号，下次遇到还是错同样的地方。😤后来用诗雨漫画笔记把背后原理搞懂，错题才真的不再重复。",
        "extra": ["#会计学习方法"],
    },
    {
        "num": 212, "account": "孙文礼",
        "title": "学了新的忘了旧的很挫败",
        "body": "学到后面章节，前面学的又忘得差不多了，感觉像在做无用功。😵后来常翻诗雨漫画笔记回顾框架，前后串起来记得更牢。",
        "extra": ["#中级会计备考"],
    },
    {
        "num": 213, "account": "晴天",
        "title": "考前突然怀疑自己复习方向",
        "body": "越接近考试越慌，怀疑自己这段时间是不是复习错了方向。😮‍💨后来对照诗雨漫画笔记的高频考点查漏补缺，心里踏实多了。",
        "extra": ["#考前冲刺"],
    },
]

def check(title, body):
    tc = len(title)
    bc = len(body.replace("\n", "").replace(" ", ""))
    ok = "✅" if tc <= 20 and bc <= 80 else f"❌(标题{tc},正文{bc}字)"
    return ok, tc, bc

def gen_txts():
    for post in POSTS:
        title = post["title"]
        body = post["body"]
        tags = [P_BASE] + post["extra"] + [P_TAIL]
        topics = " ".join(tags)
        account = f"14芳芳财会|{post['account']}"
        ok, tc, bc = check(title, body)
        print(f"  {ok} 标题{tc}字 正文{bc}字 | {account}")
        keyword = title[:8]
        fname = f"{post['num']}_{keyword}_{post['account']}.txt"
        path = os.path.join(BASE, fname)
        content = f"标题：{title}\n\n正文：\n{body}\n\n话题：{topics}\n\n时间：\n\n账号：{account}\n"
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        post["_fname"] = fname

def gen_word():
    doc = Document()
    for i, post in enumerate(POSTS):
        body = post["body"]
        tags = [P_BASE] + post["extra"] + [P_TAIL]
        kept_tags = []
        for group in tags:
            for tag in group.split():
                if tag not in STRIP_TAGS and tag not in kept_tags:
                    kept_tags.append(tag)
        p1 = doc.add_paragraph()
        r1 = p1.add_run(body)
        r1.font.size = Pt(12)
        p2 = doc.add_paragraph()
        r2 = p2.add_run(" ".join(kept_tags))
        r2.font.size = Pt(12)
        if i != len(POSTS) - 1:
            p3 = doc.add_paragraph()
            run = p3.add_run()
            run.add_break(WD_BREAK.PAGE)
    path = os.path.join(BASE, DOCNAME)
    doc.save(path)
    print(f"  ✅ Word汇总已保存：{path}")

def make_zip():
    zip_path = os.path.join(BASE, ZIPNAME)
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for post in POSTS:
            zf.write(os.path.join(BASE, post["_fname"]), post["_fname"])
        zf.write(os.path.join(BASE, DOCNAME), DOCNAME)
    return zip_path

if __name__ == "__main__":
    print("=== 生成9篇引流txt ===")
    gen_txts()
    print("\n=== 生成Word汇总 ===")
    gen_word()
    zp = make_zip()
    print(f"\n  ✅ ZIP打包完成：{zp}")
