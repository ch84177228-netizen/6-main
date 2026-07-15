#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import os, zipfile
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_BREAK

BASE = "/home/user/6-main/01-6个主号内容创作/引流小号_发布文案"
ZIPNAME = "编号20260705-01_引流账号文案_芳芳财会.zip"
DOCNAME = "引流小号_第178-186篇.docx"
P_BASE = "#会计 #中级会计 #中级会计备考 #注会cpa #财会"
P_TAIL = "#诗雨会计"
STRIP_TAGS = ["#财会", "#备考日记", "#上班族备考", "#考前冲刺"]

POSTS = [
    {
        "num": 178, "account": "沈小辉",
        "title": "一边带娃一边备考两头忙",
        "body": "白天带娃，晚上想学习，结果两头都顾不好，挺崩溃的。😮‍💨后来趁孩子睡了翻两页诗雨漫画笔记，见缝插针也能学点。",
        "extra": ["#备考日记"],
    },
    {
        "num": 179, "account": "孙青5263",
        "title": "报名后教材一直没打开心虚",
        "body": "钱交了、名报了，教材买回来一直没打开，越拖越心虚。😅后来逼自己先翻诗雨漫画笔记两页，好歹迈出第一步。",
        "extra": ["#备考焦虑"],
    },
    {
        "num": 180, "account": "橙子学会计",
        "title": "总等大块时间结果一直没有",
        "body": "总想着等哪天有一整块空闲时间好好学，结果这块时间从来没出现过。😮‍💨后来用诗雨漫画笔记利用碎片时间，反而学得更多。",
        "extra": ["#备考日记"],
    },
    {
        "num": 181, "account": "新老婆",
        "title": "学着学着手机一刷半小时没了",
        "body": "本来想学十分钟，结果手机一拿起来刷半小时，回头看书已经没心情了。😩后来把手机放远，翻诗雨漫画笔记不容易分心。",
        "extra": ["#备考日记"],
    },
    {
        "num": 182, "account": "孙玲小红书",
        "title": "做题总卡在最后一道大题",
        "body": "前面小题都能应付，一到最后的大题就卡壳，特别打击信心。😣后来靠诗雨漫画笔记把思路理顺，大题也没那么怕了。",
        "extra": ["#备考日记"],
    },
    {
        "num": 183, "account": "爹爹",
        "title": "年纪大了记性差怕考不过",
        "body": "年纪不小了，总觉得记性不如年轻人，怕怎么背都背不进去。😔后来发现诗雨漫画笔记靠图理解记，比死记硬背管用多了。",
        "extra": ["#备考焦虑"],
    },
    {
        "num": 184, "account": "张菊香",
        "title": "复习进度总跟不上老师课程",
        "body": "老师课都讲到后面了，自己复习还停在前几章，越落越多越慌。😮‍💨后来先靠诗雨漫画笔记把高频考点补上，没那么慌了。",
        "extra": ["#备考焦虑"],
    },
    {
        "num": 185, "account": "孙文礼",
        "title": "想一口气学完一章结果学到一半",
        "body": "总想一口气啃完一整章，结果学到一半就撑不住放弃了。😩后来用诗雨漫画笔记拆着看，一小节一小节反而能坚持下来。",
        "extra": ["#备考日记"],
    },
    {
        "num": 186, "account": "晴天",
        "title": "交完报名费瞬间才开始慌",
        "body": "报名费一交，瞬间意识到考试是真的了，才开始认真慌。😱赶紧翻出诗雨漫画笔记把基础过一遍，好歹不算两手空空。",
        "extra": ["#考前冲刺"],
    },
]

def check(title, body):
    tc = len(title)
    bc = len(body.replace("\n", "").replace(" ", ""))
    ok = "✅" if tc <= 20 and bc <= 80 else f"❌(标题{tc},正文{bc}字)"
    return ok, tc, bc

def gen_txts():
    for post in POSTS:
        title = post["title"]
        body = post["body"]
        topics = " ".join([P_BASE] + post["extra"] + [P_TAIL])
        account = f"14芳芳财会|{post['account']}"
        ok, tc, bc = check(title, body)
        print(f"  {ok} 标题{tc}字 正文{bc}字 | {account}")
        keyword = title[:8]
        fname = f"{post['num']}_{keyword}_{post['account']}.txt"
        path = os.path.join(BASE, fname)
        content = f"标题：{title}\n\n正文：\n{body}\n\n话题：{topics}\n\n时间：\n\n账号：{account}\n"
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        post["_fname"] = fname

def gen_word():
    doc = Document()
    for i, post in enumerate(POSTS):
        body = post["body"]
        topics_list = [P_BASE] + post["extra"] + [P_TAIL]
        kept_tags = []
        for group in topics_list:
            for tag in group.split():
                if tag not in STRIP_TAGS and tag not in kept_tags:
                    kept_tags.append(tag)
        p1 = doc.add_paragraph()
        r1 = p1.add_run(body)
        r1.font.size = Pt(12)
        p2 = doc.add_paragraph()
        r2 = p2.add_run(" ".join(kept_tags))
        r2.font.size = Pt(12)
        if i != len(POSTS) - 1:
            p3 = doc.add_paragraph()
            run = p3.add_run()
            run.add_break(WD_BREAK.PAGE)
    path = os.path.join(BASE, DOCNAME)
    doc.save(path)
    print(f"  ✅ Word汇总已保存：{path}")

def make_zip():
    zip_path = os.path.join(BASE, ZIPNAME)
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for post in POSTS:
            zf.write(os.path.join(BASE, post["_fname"]), post["_fname"])
        zf.write(os.path.join(BASE, DOCNAME), DOCNAME)
    return zip_path

if __name__ == "__main__":
    print("=== 生成9篇引流txt ===")
    gen_txts()
    print("\n=== 生成Word汇总 ===")
    gen_word()
    zp = make_zip()
    print(f"\n  ✅ ZIP打包完成：{zp}")
