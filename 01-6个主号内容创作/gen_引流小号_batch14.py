#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import os, zipfile
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_BREAK

BASE = "/home/user/6-main/01-6个主号内容创作/引流小号_发布文案"
P = "#会计 #中级会计 #中级会计备考 #注会cpa #财会 #备考日记 #上班族备考 #考前冲刺 #诗雨会计"
# Word汇总需删除的标签
DROP_TAGS = {"#财会", "#备考日记", "#上班族备考", "#考前冲刺"}

def save(path, content):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)

records = []

def write_txt(idx, title, body, account_name):
    tc = len(title)
    bc = len(body.replace("\n","").replace(" ",""))
    ok = "✅" if tc<=20 and bc<=80 else f"❌(标题{tc},正文{bc}字)"
    print(f"  {ok} 标题{tc}字 正文{bc}字 | {account_name}")
    kw = title[:8].replace("？","").replace("！","").replace("，","").replace("。","")
    fname = f"{idx}_{kw}_{account_name}.txt"
    content = f"标题：{title}\n\n正文：\n{body}\n\n话题：{P}\n\n时间：\n\n账号：14芳芳财会|{account_name}\n"
    save(os.path.join(BASE, fname), content)
    records.append((idx, body))

def make_zip(start, end):
    zip_path = os.path.join(BASE, f"引流小号_第{start}-{end}篇.zip")
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for i in range(start, end+1):
            for fname in os.listdir(BASE):
                if fname.startswith(f"{i}_") and fname.endswith(".txt"):
                    zf.write(os.path.join(BASE, fname), fname)
        # 加入配套Word
        wname = f"引流小号_第{start}-{end}篇.docx"
        if os.path.exists(os.path.join(BASE, wname)):
            zf.write(os.path.join(BASE, wname), wname)
    return zip_path

def gen_word(start, end):
    doc = Document()
    first = True
    for idx, body in records:
        if not first:
            p = doc.add_paragraph()
            p.add_run().add_break(WD_BREAK.PAGE)
        h = doc.add_paragraph()
        hr = h.add_run(f"第{idx}篇")
        hr.bold = True
        hr.font.size = Pt(14)
        for line in body.strip().split("\n"):
            p = doc.add_paragraph()
            r = p.add_run(line)
            r.font.size = Pt(12)
        tags = [t for t in P.split() if t not in DROP_TAGS]
        p = doc.add_paragraph()
        r = p.add_run(" ".join(tags))
        r.font.size = Pt(12)
        first = False
    out = os.path.join(BASE, f"引流小号_第{start}-{end}篇.docx")
    doc.save(out)
    print(f"  ✅ 汇总Word：{out}")

articles = [
    (124, "考前69天才发现进度严重落后",
     "倒计时69天，翻开计划表发现自己连一半都没学完，整个人都焦虑了。😰后来用诗雨漫画笔记抓重点章节，效率高了不少，落后的进度慢慢追了一点回来。",
     "沈小辉"),
    (125, "长投处置那块怎么学都绕不明白",
     "长期股权投资处置、转换那几种情况，看一遍忘一遍，脑子一团乱。🤯翻了诗雨漫画笔记里的对比图，几种转换一目了然，终于理顺了。",
     "孙青5263"),
    (126, "白天上班晚上看书眼睛都花了",
     "上了一天班，晚上还要硬撑着看书，看着看着眼睛就花了，效率很低。😣换成漫画版笔记，图多字少，看着没那么累，能多坚持一会儿。",
     "橙子学会计"),
    (127, "所得税递延那块是我的噩梦",
     "递延所得税资产负债，账面计税基础比来比去，每次都搞反方向。😵看了诗雨漫画笔记的图解，资产高交税资产低省税，一下记住了。",
     "新老婆"),
    (128, "刷题总是同一个坑反复踩",
     "做题错来错去就那几个知识点，明明对过答案，下次还是错。😤后来发现是没真懂。配合诗雨漫画笔记把考点重新过了一遍，错得少了。",
     "孙玲小红书"),
    (129, "看着别人晒进度我更慌了",
     "备考群里别人都说二刷三刷了，我才一刷，越看越慌干脆退了群。😮‍💨专心用诗雨漫画笔记按自己节奏走，反而踏实多了。",
     "爹爹"),
    (130, "记了满满一本笔记却没用上",
     "手抄了厚厚一本笔记，结果做题时一条都想不起来，白费功夫。😭后来用诗雨漫画笔记，图像记忆比抄写有用，关键考点真记住了。",
     "张菊香"),
    (131, "下决心早起背书结果起不来",
     "立flag要早起一小时背书，结果闹钟响了又睡，计划全泡汤。😅改成通勤时间用诗雨漫画笔记翻几页，碎片时间反而坚持下来了。",
     "孙文礼"),
    (132, "越接近考试越怕自己考不过",
     "还剩两个多月，越想越怕万一考不过怎么办，压力大到失眠。😔后来告诉自己稳扎稳打，用诗雨漫画笔记一天搞懂一个点，心里踏实些了。",
     "晴天"),
]

if __name__ == "__main__":
    print("=== 生成引流小号第124-132篇 ===")
    for idx, title, body, account in articles:
        write_txt(idx, title, body, account)
    gen_word(124, 132)
    zp = make_zip(124, 132)
    print(f"\n  ✅ ZIP打包完成：{zp}")
