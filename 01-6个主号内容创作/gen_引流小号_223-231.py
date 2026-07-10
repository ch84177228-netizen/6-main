#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import os, zipfile
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_BREAK

BASE = "/home/user/6-main/01-6个主号内容创作/引流小号_发布文案"
ZIPNAME = "编号20260709-08_引流账号文案_芳芳财会.zip"
DOCNAME = "引流小号_第223-231篇.docx"
P_BASE = "#会计 #中级会计 #中级会计备考 #注会cpa #财会"
P_TAIL = "#诗雨会计"
# 本批要求：Word 里的话题固定保留 5 个
WORD_TAGS = ["#会计", "#中级会计", "#中级会计备考", "#注会cpa", "#诗雨会计"]

POSTS = [
    {
        "num": 223, "account": "沈小辉",
        "title": "报了名却总把学习往后推",
        "body": "名报了、书买了，可每天都想着明天再学，一拖就是好几天。😮‍💨后来把诗雨漫画笔记放手边，翻开就学，少了很多纠结。",
        "extra": ["#上班族备考"],
    },
    {
        "num": 224, "account": "孙青5263",
        "title": "零基础翻开教材完全看不懂",
        "body": "第一次接触会计，翻开教材满眼陌生名词，完全看不懂。😵后来先跟着诗雨漫画笔记打基础，一点点入门，没那么慌了。",
        "extra": ["#零基础备考"],
    },
    {
        "num": 225, "account": "橙子学会计",
        "title": "分录记了一堆一到题就乱",
        "body": "会计分录背了一堆，一到综合题就分不清借贷、记乱了。😩后来靠诗雨漫画笔记把分录逻辑理顺，做题才没那么慌。",
        "extra": ["#中级会计实务"],
    },
    {
        "num": 226, "account": "新老婆",
        "title": "看书容易犯困坚持不了多久",
        "body": "一坐下来看书没多久就开始犯困，坚持不了太长时间。😴后来换成诗雨漫画笔记，图多字少看着轻松，反而能多学一会儿。",
        "extra": ["#会计学习方法"],
    },
    {
        "num": 227, "account": "孙玲小红书",
        "title": "总想等状态好再开始结果一直没开始",
        "body": "总想着等状态调好、准备充分再认真学，结果拖到现在还没真正开始。😅后来先翻两页诗雨漫画笔记，开了头就没那么难。",
        "extra": ["#会计学习方法"],
    },
    {
        "num": 228, "account": "爹爹",
        "title": "年纪大了怕自己记不住",
        "body": "年纪不小了，总担心记性不如年轻人、怎么背都背不牢。😔后来发现诗雨漫画笔记靠图理解记，比死记硬背轻松多了。",
        "extra": ["#上班族备考"],
    },
    {
        "num": 229, "account": "张菊香",
        "title": "错题改了下次还是错同一处",
        "body": "错题认真改了、也标了记号，下次遇到还是错在同一个地方。😤后来用诗雨漫画笔记把原理搞懂，错题才真的不再犯。",
        "extra": ["#中级会计备考"],
    },
    {
        "num": 230, "account": "孙文礼",
        "title": "学到后面就忘了前面",
        "body": "学到后面章节，前面学的又忘得差不多，感觉像白学。😵后来常翻诗雨漫画笔记回顾框架，前后串起来记得更牢。",
        "extra": ["#中级会计备考"],
    },
    {
        "num": 231, "account": "晴天",
        "title": "越接近考试越乱越没底",
        "body": "越接近考试心里越乱，越觉得自己没复习到位、心里没底。😮‍💨后来靠诗雨漫画笔记把高频考点一遍遍过，慢慢踏实了。",
        "extra": ["#考前冲刺"],
    },
]

def check(title, body):
    tc = len(title)
    bc = len(body.replace("\n", "").replace(" ", ""))
    ok = "✅" if tc <= 20 and bc <= 80 else f"❌(标题{tc},正文{bc}字)"
    return ok, tc, bc

def gen_txts():
    for post in POSTS:
        title = post["title"]
        body = post["body"]
        tags = [P_BASE] + post["extra"] + [P_TAIL]
        topics = " ".join(tags)
        account = f"14芳芳财会|{post['account']}"
        ok, tc, bc = check(title, body)
        print(f"  {ok} 标题{tc}字 正文{bc}字 | {account}")
        keyword = title[:8]
        fname = f"{post['num']}_{keyword}_{post['account']}.txt"
        path = os.path.join(BASE, fname)
        content = f"标题：{title}\n\n正文：\n{body}\n\n话题：{topics}\n\n时间：\n\n账号：{account}\n"
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        post["_fname"] = fname

def gen_word():
    doc = Document()
    tag_line = " ".join(WORD_TAGS)  # 固定5个话题
    for i, post in enumerate(POSTS):
        body = post["body"]
        p1 = doc.add_paragraph()
        r1 = p1.add_run(body)
        r1.font.size = Pt(12)
        p2 = doc.add_paragraph()
        r2 = p2.add_run(tag_line)
        r2.font.size = Pt(12)
        if i != len(POSTS) - 1:
            p3 = doc.add_paragraph()
            run = p3.add_run()
            run.add_break(WD_BREAK.PAGE)
    path = os.path.join(BASE, DOCNAME)
    doc.save(path)
    print(f"  ✅ Word汇总已保存（每篇话题固定5个）：{path}")

def make_zip():
    zip_path = os.path.join(BASE, ZIPNAME)
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for post in POSTS:
            zf.write(os.path.join(BASE, post["_fname"]), post["_fname"])
        zf.write(os.path.join(BASE, DOCNAME), DOCNAME)
    return zip_path

if __name__ == "__main__":
    print("=== 生成9篇引流txt ===")
    gen_txts()
    print("\n=== 生成Word汇总 ===")
    gen_word()
    zp = make_zip()
    print(f"\n  ✅ ZIP打包完成：{zp}")
