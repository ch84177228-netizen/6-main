#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

OUT = "/home/user/6-main/12账号分组与引流号对照表.xlsx"

HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="center", horizontal="left")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")

def style_sheet(ws, header_row, col_widths, center_cols=None):
    center_cols = center_cols or []
    for cell in ws[header_row]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
        for cell in row:
            cell.border = BORDER
            cell.alignment = CENTER if cell.column_letter in center_cols else WRAP
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1).coordinate

wb = Workbook()

# ===== 表1：12账号写作风格 =====
ws1 = wb.active
ws1.title = "12账号写作风格"
ws1.append(["账号", "定位/风格", "标题示例（存货主题）"])
rows1 = [
    ("①主号诗雨学姐", "甄嬛传故事风格——用后宫人物/场景类比知识点，需有具体数字和完整步骤", "甄嬛传版存货库房管理"),
    ("②副号1学生视角", "蜡笔小新版生活故事——用真实生活场景（广志超市/家庭消费）带入知识点，篇幅较短、故事轻快", "广志超市学存货计价法"),
    ("③副号2干货速记", "纯干货清单体，公式+分录直接罗列，适合快速记忆", "存货核算速记清单"),
    ("④副号3避坑吐槽", "错误写法vs正确写法对比体，专挑历年最容易踩的坑", "存货跌价准备最易错"),
    ("⑤副号4纯考点总结", "考点清单体，按定义→分类→公式→规则结构化整理", "存货核心考点整理"),
    ("⑥副号5考前冲刺", "开篇带距考试还有X天倒计时（按当天日期自动计算），营造紧迫感", "距中级还有65天存货冲刺"),
    ("⑦西游记故事号", "用师徒四人经历类比知识点", "取经队伍干粮存货管理"),
    ("⑧小新故事号", "小新家庭生活场景（区别于②，场景更家庭化）", "小新家零食过期踩坑"),
    ("⑨治愈猫咪插图号", "猫咪拟人化喵～语气，配图解风格讲解", "猫咪图解存货跌价准备"),
    ("⑩西游记故事号2", "八戒踩坑视角，专讲错误示范", "八戒囤货变质踩坑存货"),
    ("⑪小新故事号2", "小新图解对比体（区别于⑧，侧重方法对比）", "小新图解发出计价方法"),
    ("⑫会计科普考点全览版", "背诵口诀+速查清单，收尾总结全篇", "存货考点速查背诵"),
]
for r in rows1:
    ws1.append(list(r))
style_sheet(ws1, 1, [20, 55, 30])

# ===== 表2：12组账号对照表 =====
ws2 = wb.create_sheet("12组账号对照表")
ws2.append(["组", "抖音", "小红书", "视频号", "公众号"])
rows2 = [
    ("第1组", "07抖音发作品|01蓝色", "06红书店铺|粉色手机", "04视频号|01主账号黄色手机", "13公众号|01蓝色密码"),
    ("第2组", "07抖音发作品|02粉色西瓜", "06红书店铺|黄色手机", "04视频号|02诗雨vip", "13公众号|02VIPSun"),
    ("第3组", "07抖音发作品|03黄色", "06红书店铺|蓝色ipad张菊芳", "04视频号|03蓝色手机", "13公众号|03黄色手机"),
    ("第4组", "07抖音发作品|04粉色糖糖", "06红书店铺|孙文新134 已实名", "04视频号|04孙文礼", "13公众号|04吴学安"),
    ("第5组", "07抖音发作品|05吴学安手机号", "06红书店铺|孙文新 小号店铺号", "04视频号|05胡芝兰", "—"),
    ("第6组", "07抖音发作品|06蓝色", "06红书店铺|姗姗已实名", "04视频号|06吴学安", "—"),
    ("第7组", "07抖音发作品|07会计漫画", "06红书店铺|孙咏美", "—", "—"),
    ("第8组", "07抖音发作品|08文新", "06红书店铺|姗姗 店铺大号", "—", "—"),
    ("第9组", "07抖音发作品|09专业号", "06红书店铺|西瓜学姐", "—", "—"),
    ("第10组", "—（999学霸已停用）", "06红书店铺|会计漫画", "—", "—"),
    ("第11组", "—", "06红书店铺|樱桃学姐", "—", "—"),
    ("第12组", "—", "06红书店铺|胡志兰手机", "—", "—"),
]
for r in rows2:
    ws2.append(list(r))
style_sheet(ws2, 1, [10, 26, 26, 26, 22], center_cols=["A"])

ws2.append([])
note_row = ws2.max_row + 1
ws2.cell(row=note_row, column=1, value="文件名前缀 / 账号字段格式规则（最高优先级，绝对不能错）：")
ws2.cell(row=note_row, column=1).font = Font(bold=True)
rule_rows = [
    ("平台", "文件名前缀", "账号字段格式"),
    ("抖音", "07_抖音_账号名.txt", "07抖音发作品|账号名"),
    ("小红书", "06_小红书_账号名.txt", "06红书店铺|账号名"),
    ("视频号", "04_视频号_账号名.txt", "04视频号|账号名"),
    ("公众号", "13_公众号_账号名.txt", "13公众号|账号名"),
]
start = note_row + 1
for i, r in enumerate(rule_rows):
    for j, val in enumerate(r, start=1):
        c = ws2.cell(row=start + i, column=j, value=val)
        c.border = BORDER
        c.alignment = CENTER if j == 1 else WRAP
        if i == 0:
            c.font = HEADER_FONT
            c.fill = HEADER_FILL

# ===== 表3：引流小号9个固定账号 =====
ws3 = wb.create_sheet("引流小号9账号")
ws3.append(["序号", "账号（账号字段：14芳芳财会|账号名）"])
rows3 = [
    (1, "14芳芳财会|沈小辉"),
    (2, "14芳芳财会|孙青5263"),
    (3, "14芳芳财会|橙子学会计"),
    (4, "14芳芳财会|新老婆"),
    (5, "14芳芳财会|孙玲小红书"),
    (6, "14芳芳财会|爹爹"),
    (7, "14芳芳财会|张菊香"),
    (8, "14芳芳财会|孙文礼"),
    (9, "14芳芳财会|晴天"),
]
for r in rows3:
    ws3.append(list(r))
style_sheet(ws3, 1, [10, 40], center_cols=["A"])

ws3.append([])
note_row3 = ws3.max_row + 1
ws3.cell(row=note_row3, column=1, value="其他规则要点：")
ws3.cell(row=note_row3, column=1).font = Font(bold=True)
notes = [
    "文件命名格式：序号_标题关键词_账号名.txt（序号为历史累计篇号，非1-9）",
    "固定生成9篇小红书文案txt + 1个Word汇总",
    "Word汇总只保留正文+话题，去掉标题/时间/账号字段",
    "Word需去除 #财会 #备考日记 #上班族备考 #考前冲刺 标签（txt保留）",
    "生成前需核对历史文案，标题正文不得重复",
    "正文需软植入\"诗雨漫画笔记\"，不能硬广（不可出现私信我/关注我等话术）",
    "默认不在对话中显示文案内容，直接生成文件打包发ZIP",
]
for i, n in enumerate(notes):
    ws3.cell(row=note_row3 + 1 + i, column=1, value=f"· {n}")

wb.save(OUT)
print(f"✅ 已生成：{OUT}")
